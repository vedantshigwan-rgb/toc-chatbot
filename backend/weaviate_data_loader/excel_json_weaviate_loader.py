from __future__ import annotations

import json
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

from openpyxl import load_workbook

from config.settings import OPENAI_API_KEY


# ============================================================
# CONFIG
# ============================================================

EXCEL_FILE = Path(
    r"D:\AI gen\TOC_backend_primary_only\data\JSON METADATA.xlsx"
)

JSON_FILE = Path(
    r"D:\AI gen\backend_v4\data\column_metadata.json"
)

WEAVIATE_URL = "http://localhost:8080"

CLASS_NAME = "ColumnMetadata"

ROLE_MASTER_SHEET = "ROLE_MASTER"

BATCH_SIZE = 50


# ============================================================
# REPORT -> MODULE ID -> TABLE MAPPING
# ============================================================

REPORT_TABLE_MAP: Dict[str, Tuple[str, str]] = {

    "Inventory Report": (
        "01",
        "FULLKIT_ITEMWISE_INV",
    ),

    "Buyer Wise Inventory Report": (
        "02",
        "KPC_INVENTORY_BUYER_REP",
    ),

    "Order Report": (
        "03",
        "KPC_FULL_KIT_OA_GROUPING",
    ),

    "BOM Report": (
        "04",
        "KPC_FULL_KIT_BOM_EXPLODE",
    ),

    "Docket Report": (
        "05",
        "KPC_FK_DOCKET_IDLE_REPORT",
    ),

    "Purchase Order Report": (
        "06",
        "KPC_PO_BOARD",
    ),

    "Docket Employee Report": (
        "07",
        "KPC_FULL_KIT_DOCKET_EMP",
    ),

    "On Time Delivery Index Report": (
        "08",
        "KPC_FULL_KIT_ONTIME_DELIVERY_INDEX",
    ),

    "OA Trolley Report": (
        "09",
        "OA_TROLLEY_STATUS",
    ),

    "Buyerwise OA Trolley Report": (
        "10",
        "BUYERWISE_OA_TROLLEY_STATUS",
    ),

    "POB COGS Report": (
        "11",
        "POB_COGS",
    ),

    "POB COGS Spares Report": (
        "12",
        "POB_COGS_SPARES",
    ),
}


# ============================================================
# CUSTOM ID
#
# Example:
#
# 01.001.021
#
# 01  = Report / Module
# 001 = Table inside report
# 021 = Column inside table
# ============================================================

CUSTOM_ID_PATTERN = re.compile(
    r"^\d{2}\.\d{3}\.\d{3}$"
)


# ============================================================
# REQUIRED COLUMN METADATA PROPERTIES
# ============================================================

REQUIRED_PROPERTIES = {
    "text",
    "columnName",
    "tableName",
    "module",
    "dataType",
    "description",
    "distinctValues",
    "units",
    "used",
    "filterable",
    "roleAccess",
}


# ============================================================
# BASIC HELPERS
# ============================================================

def normalize_header(value: Any) -> str:
    """
    Normalize Excel header names so:

        'JSON object'
        ' JSON Object '
        'json   object'

    all become:

        'json object'
    """

    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip(),
    ).lower()


def get_headers(ws) -> Dict[str, int]:
    """
    Returns normalized header -> Excel column number.
    """

    return {
        normalize_header(cell.value): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


# ============================================================
# ROLE MASTER
# ============================================================

def read_role_master(
    workbook,
) -> Tuple[Set[str], Dict[str, str]]:

    if ROLE_MASTER_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"Missing sheet: {ROLE_MASTER_SHEET}"
        )

    ws = workbook[ROLE_MASTER_SHEET]

    headers = get_headers(ws)

    role_id_col = headers.get("role_id")
    role_name_col = headers.get("role_name")

    if not role_id_col or not role_name_col:
        raise ValueError(
            "ROLE_MASTER must contain "
            "'role_id' and 'role_name' columns."
        )

    role_ids: Set[str] = set()

    role_map: Dict[str, str] = {}

    for row in range(2, ws.max_row + 1):

        role_id = ws.cell(
            row,
            role_id_col,
        ).value

        role_name = ws.cell(
            row,
            role_name_col,
        ).value

        # Ignore completely blank rows.
        if role_id is None and role_name is None:
            continue

        role_id = (
            ""
            if role_id is None
            else str(role_id).strip()
        )

        role_name = (
            ""
            if role_name is None
            else str(role_name).strip()
        )

        if not role_id or not role_name:
            raise ValueError(
                f"Invalid ROLE_MASTER row {row}: "
                f"role_id={role_id!r}, "
                f"role_name={role_name!r}"
            )

        if role_id in role_ids:
            raise ValueError(
                f"Duplicate role_id in ROLE_MASTER: "
                f"{role_id}"
            )

        role_ids.add(role_id)

        role_map[role_id] = role_name

    if not role_ids:
        raise ValueError(
            "ROLE_MASTER contains no roles."
        )

    print(
        f"[OK] ROLE_MASTER loaded: "
        f"{len(role_ids)} roles"
    )

    return role_ids, role_map


# ============================================================
# VALIDATE ONE METADATA OBJECT
# ============================================================

def validate_object(
    obj: Dict[str, Any],
    *,
    sheet_name: str,
    excel_row: int,
    module_id: str,
    table_name: str,
    column_number: int,
    excel_column_name: str,
    excel_data_type: str | None,
    allowed_roles: Set[str],
) -> Dict[str, Any]:

    location = (
        f"{sheet_name} row {excel_row}"
    )

    # --------------------------------------------------------
    # JSON object itself
    # --------------------------------------------------------

    if not isinstance(obj, dict):
        raise ValueError(
            f"{location}: JSON must be an object"
        )

    # --------------------------------------------------------
    # custom_id
    # --------------------------------------------------------

    custom_id = obj.get("custom_id")

    if not isinstance(custom_id, str):
        raise ValueError(
            f"{location}: custom_id must be string"
        )

    custom_id = custom_id.strip()

    if not CUSTOM_ID_PATTERN.fullmatch(custom_id):
        raise ValueError(
            f"{location}: invalid custom_id "
            f"{custom_id!r}"
        )

    expected_id = (
        f"{module_id}.001."
        f"{column_number:03d}"
    )

    if custom_id != expected_id:
        raise ValueError(
            f"{location}: custom_id mismatch. "
            f"Found={custom_id}, "
            f"Expected={expected_id}"
        )

    # --------------------------------------------------------
    # class
    # --------------------------------------------------------

    object_class = obj.get("class")

    if object_class != CLASS_NAME:
        raise ValueError(
            f"{location}: "
            f"class must be {CLASS_NAME!r}, "
            f"found {object_class!r}"
        )

    # --------------------------------------------------------
    # properties
    # --------------------------------------------------------

    properties = obj.get("properties")

    if not isinstance(properties, dict):
        raise ValueError(
            f"{location}: "
            f"properties must be an object"
        )

    missing_properties = (
        REQUIRED_PROPERTIES
        - set(properties.keys())
    )

    if missing_properties:
        raise ValueError(
            f"{location}: "
            f"missing properties "
            f"{sorted(missing_properties)}"
        )

    # --------------------------------------------------------
    # Module / report
    # --------------------------------------------------------

    module = str(
        properties.get("module", "")
    ).strip()

    if module != sheet_name:
        raise ValueError(
            f"{location}: "
            f"module mismatch. "
            f"Found={module!r}, "
            f"Expected={sheet_name!r}"
        )

    # --------------------------------------------------------
    # Table
    # --------------------------------------------------------

    actual_table_name = str(
        properties.get(
            "tableName",
            "",
        )
    ).strip()

    if actual_table_name != table_name:
        raise ValueError(
            f"{location}: "
            f"tableName mismatch. "
            f"Found={actual_table_name!r}, "
            f"Expected={table_name!r}"
        )

    # --------------------------------------------------------
    # Column
    # --------------------------------------------------------

    json_column_name = str(
        properties.get(
            "columnName",
            "",
        )
    ).strip()

    if json_column_name != excel_column_name:
        raise ValueError(
            f"{location}: "
            f"columnName mismatch. "
            f"Excel={excel_column_name!r}, "
            f"JSON={json_column_name!r}"
        )

    # --------------------------------------------------------
    # Data type
    # --------------------------------------------------------

    if excel_data_type is not None:

        json_data_type = str(
            properties.get(
                "dataType",
                "",
            )
        ).strip()

        if json_data_type != excel_data_type:
            raise ValueError(
                f"{location}: "
                f"dataType mismatch. "
                f"Excel={excel_data_type!r}, "
                f"JSON={json_data_type!r}"
            )

    # --------------------------------------------------------
    # distinctValues
    # --------------------------------------------------------

    if not isinstance(
        properties.get("distinctValues"),
        list,
    ):
        raise ValueError(
            f"{location}: "
            f"distinctValues must be a list"
        )

    # --------------------------------------------------------
    # units
    # --------------------------------------------------------

    if not isinstance(
        properties.get("units"),
        list,
    ):
        raise ValueError(
            f"{location}: units must be a list"
        )

    # --------------------------------------------------------
    # used
    # --------------------------------------------------------

    used = properties.get("used")

    if used not in (0, 1):
        raise ValueError(
            f"{location}: "
            f"used must be 0 or 1"
        )

    # --------------------------------------------------------
    # filterable
    # --------------------------------------------------------

    filterable = properties.get(
        "filterable"
    )

    if not isinstance(filterable, bool):
        raise ValueError(
            f"{location}: "
            f"filterable must be boolean"
        )

    # --------------------------------------------------------
    # roleAccess
    # --------------------------------------------------------

    role_access = properties.get(
        "roleAccess"
    )

    if not isinstance(role_access, list):
        raise ValueError(
            f"{location}: "
            f"roleAccess must be a list"
        )

    unknown_roles = [
        role
        for role in role_access
        if role not in allowed_roles
    ]

    if unknown_roles:
        raise ValueError(
            f"{location}: "
            f"unknown role IDs "
            f"{unknown_roles}"
        )

    if len(role_access) != len(
        set(role_access)
    ):
        raise ValueError(
            f"{location}: "
            f"duplicate role IDs in roleAccess"
        )

    # --------------------------------------------------------
    # IMPORTANT
    #
    # custom_id exists at top level in our JSON file,
    # but Weaviate only stores/searches properties.
    #
    # Therefore copy it inside properties.
    # --------------------------------------------------------

    properties = dict(properties)

    properties["custom_id"] = custom_id

    return {
        "custom_id": custom_id,
        "class": CLASS_NAME,
        "properties": properties,
    }


# ============================================================
# STEP 1
#
# EXCEL -> ONE VALIDATED JSON FILE
# ============================================================

def excel_to_json(
    excel_path: Path,
    json_path: Path,
) -> List[Dict[str, Any]]:

    # --------------------------------------------------------
    # File exists
    # --------------------------------------------------------

    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: "
            f"{excel_path}"
        )

    print(
        f"[INFO] Reading workbook: "
        f"{excel_path}"
    )

    workbook = load_workbook(
        excel_path,
        data_only=False,
    )

    # --------------------------------------------------------
    # Roles
    # --------------------------------------------------------

    allowed_roles, role_map = (
        read_role_master(workbook)
    )

    # --------------------------------------------------------
    # Metadata sheets
    # --------------------------------------------------------

    metadata_sheets = [
        name
        for name in workbook.sheetnames
        if name != ROLE_MASTER_SHEET
    ]

    expected_sheets = set(
        REPORT_TABLE_MAP.keys()
    )

    actual_sheets = set(
        metadata_sheets
    )

    missing_sheets = (
        expected_sheets
        - actual_sheets
    )

    unexpected_sheets = (
        actual_sheets
        - expected_sheets
    )

    if missing_sheets or unexpected_sheets:

        raise ValueError(
            "\nReport sheet mismatch.\n"
            f"Missing: "
            f"{sorted(missing_sheets)}\n"
            f"Unexpected: "
            f"{sorted(unexpected_sheets)}"
        )

    # --------------------------------------------------------
    # Storage
    # --------------------------------------------------------

    objects: List[Dict[str, Any]] = []

    seen_ids: Set[str] = set()

    sheet_counts: Dict[str, int] = {}

    # --------------------------------------------------------
    # Process each report
    # --------------------------------------------------------

    for sheet_name in metadata_sheets:

        ws = workbook[sheet_name]

        module_id, table_name = (
            REPORT_TABLE_MAP[
                sheet_name
            ]
        )

        headers = get_headers(ws)

        json_col = headers.get(
            "json object"
        )

        column_name_col = headers.get(
            "column name"
        )

        data_type_col = headers.get(
            "datatype"
        )

        if (
            not json_col
            or not column_name_col
            or not data_type_col
        ):
            raise ValueError(
                f"{sheet_name}: "
                f"required headers missing. "
                f"Need JSON object, "
                f"Column Name and DataType."
            )

        column_number = 0

        # ----------------------------------------------------
        # Rows
        # ----------------------------------------------------

        for excel_row in range(
            2,
            ws.max_row + 1,
        ):

            raw_json = ws.cell(
                excel_row,
                json_col,
            ).value

            # Ignore rows without metadata JSON.
            if (
                raw_json is None
                or not str(
                    raw_json
                ).strip()
            ):
                continue

            column_number += 1

            # ------------------------------------------------
            # Parse JSON
            # ------------------------------------------------

            try:

                parsed = json.loads(
                    str(raw_json)
                )

            except json.JSONDecodeError as exc:

                raise ValueError(
                    f"Invalid JSON: "
                    f"{sheet_name} "
                    f"row {excel_row}: "
                    f"{exc}"
                ) from exc

            # ------------------------------------------------
            # Excel column name
            # ------------------------------------------------

            raw_column_name = ws.cell(
                excel_row,
                column_name_col,
            ).value

            if raw_column_name is None:
                raise ValueError(
                    f"{sheet_name} "
                    f"row {excel_row}: "
                    f"Column Name is empty"
                )

            excel_column_name = str(
                raw_column_name
            ).strip()

            # ------------------------------------------------
            # Excel data type
            # ------------------------------------------------

            raw_data_type = ws.cell(
                excel_row,
                data_type_col,
            ).value

            excel_data_type = (
                None
                if raw_data_type is None
                else str(
                    raw_data_type
                ).strip()
            )

            # ------------------------------------------------
            # Validate
            # ------------------------------------------------

            obj = validate_object(
                parsed,
                sheet_name=sheet_name,
                excel_row=excel_row,
                module_id=module_id,
                table_name=table_name,
                column_number=column_number,
                excel_column_name=(
                    excel_column_name
                ),
                excel_data_type=(
                    excel_data_type
                ),
                allowed_roles=(
                    allowed_roles
                ),
            )

            custom_id = obj[
                "custom_id"
            ]

            # ------------------------------------------------
            # Duplicate business ID check
            # ------------------------------------------------

            if custom_id in seen_ids:
                raise ValueError(
                    f"Duplicate custom_id: "
                    f"{custom_id}"
                )

            seen_ids.add(custom_id)

            objects.append(obj)

        sheet_counts[
            sheet_name
        ] = column_number

        print(
            f"[OK] "
            f"{sheet_name}: "
            f"{column_number} objects"
        )

    # ========================================================
    # IMPORTANT:
    #
    # NO HARDCODED 700 / 701 COUNT.
    #
    # Any number is accepted as long as at least one
    # valid metadata object exists.
    # ========================================================

    if not objects:
        raise ValueError(
            "No metadata objects were found "
            "in the workbook."
        )

    # --------------------------------------------------------
    # Final duplicate check
    # --------------------------------------------------------

    all_ids = [
        obj["custom_id"]
        for obj in objects
    ]

    if len(all_ids) != len(
        set(all_ids)
    ):
        raise ValueError(
            "Duplicate custom_id values "
            "detected before JSON write."
        )

    # --------------------------------------------------------
    # JSON payload
    # --------------------------------------------------------

    payload = {

        "class": CLASS_NAME,

        "object_count": len(objects),

        "sheet_counts": (
            sheet_counts
        ),

        "role_master": (
            role_map
        ),

        "objects": objects,
    }

    # --------------------------------------------------------
    # Write JSON
    # --------------------------------------------------------

    json_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with json_path.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            payload,
            file,
            indent=2,
            ensure_ascii=False,
        )

        file.write("\n")

    # --------------------------------------------------------
    # Re-open JSON from disk
    #
    # This confirms we are going to Weaviate from a
    # genuinely valid JSON file.
    # --------------------------------------------------------

    with json_path.open(
        "r",
        encoding="utf-8",
    ) as file:

        written_payload = json.load(
            file
        )

    written_objects = (
        written_payload.get(
            "objects"
        )
    )

    if not isinstance(
        written_objects,
        list,
    ):
        raise ValueError(
            "Generated JSON verification failed: "
            "'objects' is not a list."
        )

    if len(written_objects) != len(
        objects
    ):
        raise ValueError(
            "Generated JSON verification failed: "
            "object count changed after writing."
        )

    written_ids = [
        obj.get("custom_id")
        for obj in written_objects
    ]

    if len(written_ids) != len(
        set(written_ids)
    ):
        raise ValueError(
            "Generated JSON contains "
            "duplicate custom_id values."
        )

    # --------------------------------------------------------
    # Success
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print(
        "[OK] EXCEL -> JSON COMPLETE"
    )
    print("=" * 70)

    print(
        f"[OK] JSON created: "
        f"{json_path.resolve()}"
    )

    print(
        f"[OK] Total objects: "
        f"{len(objects)}"
    )

    print()

    print(
        "Objects per report:"
    )

    for (
        sheet_name,
        count,
    ) in sheet_counts.items():

        print(
            f"  {sheet_name}: "
            f"{count}"
        )

    return objects


# ============================================================
# WEAVIATE CONNECTION
# ============================================================

def connect_to_weaviate():

    import weaviate

    # Same connection style as your current backend.
    client = weaviate.Client(

        url=WEAVIATE_URL,

        timeout_config=(
            5,
            60,
        ),

        additional_headers={
            "X-OpenAI-Api-Key":
                OPENAI_API_KEY,
        },
    )

    return client


# ============================================================
# ENSURE custom_id EXISTS AS WEAVIATE PROPERTY
# ============================================================

def ensure_custom_id_property(
    client,
) -> None:

    schema = client.schema.get()

    class_schema = next(
        (
            class_item
            for class_item
            in schema.get(
                "classes",
                [],
            )
            if class_item.get(
                "class"
            ) == CLASS_NAME
        ),
        None,
    )

    # --------------------------------------------------------
    # If ColumnMetadata doesn't exist, create it!
    # --------------------------------------------------------
    if class_schema is None:
        print(f"[SCHEMA] {CLASS_NAME} class missing, creating it...")
        schema_obj = {
            "class": CLASS_NAME,
            "description": "Metadata for Excel columns (auto-created)",
            "properties": [
                {
                    "name": "custom_id",
                    "dataType": ["text"],
                    "tokenization": "field",
                    "indexFilterable": True,
                    "indexSearchable": True,
                    "moduleConfig": { "text2vec-openai": { "skip": True, "vectorizePropertyName": False } }
                },
                {
                    "name": "columnName", "dataType": ["text"]
                },
                {
                    "name": "tableName", "dataType": ["text"]
                },
                {
                    "name": "module", "dataType": ["text"]
                },
                {
                    "name": "dataType", "dataType": ["text"]
                },
                {
                    "name": "description", "dataType": ["text"]
                },
                {
                    "name": "distinctValues", "dataType": ["text[]"]
                },
                {
                    "name": "units", "dataType": ["text[]"]
                },
                {
                    "name": "used", "dataType": ["int"]
                },
                {
                    "name": "filterable", "dataType": ["boolean"]
                },
                {
                    "name": "roleAccess", "dataType": ["text[]"]
                },
                {
                    "name": "text", "dataType": ["text"]
                }
            ],
            "moduleConfig": { "text2vec-openai": { "skip": True } }
        }
        client.schema.create_class(schema_obj)
        print(f"[OK] {CLASS_NAME} class created.")
        schema = client.schema.get()
        class_schema = next((class_item for class_item in schema.get("classes", []) if class_item.get("class") == CLASS_NAME), None)

    existing_properties = {
        prop.get("name"): prop for prop in class_schema.get("properties", [])
    }

    # --------------------------------------------------------
    # Ensure custom_id property exists (should always after above, but safe)
    # --------------------------------------------------------
    if "custom_id" not in existing_properties:
        print("[SCHEMA] Adding custom_id to ColumnMetadata...")
        client.schema.property.create(
            CLASS_NAME,
            {
                "name": "custom_id",
                "dataType": ["text"],
                "tokenization": "field",
                "indexFilterable": True,
                "indexSearchable": True,
                "moduleConfig": { "text2vec-openai": { "skip": True, "vectorizePropertyName": False } },
            },
        )
        print("[OK] custom_id property created.")
    else:
        print("[OK] custom_id property already exists.")


# ============================================================
# STEP 2
#
# JSON FILE -> WEAVIATE DOCKER
# ============================================================

def load_json_to_weaviate(
    json_path: Path,
) -> None:

    # --------------------------------------------------------
    # IMPORTANT
    #
    # Read AGAIN from JSON file.
    #
    # We do NOT load Weaviate directly from Excel memory.
    # --------------------------------------------------------

    if not json_path.exists():
        raise FileNotFoundError(
            f"JSON file not found: "
            f"{json_path}"
        )

    with json_path.open(
        "r",
        encoding="utf-8",
    ) as file:

        payload = json.load(
            file
        )

    # --------------------------------------------------------
    # Validate payload
    # --------------------------------------------------------

    objects = payload.get(
        "objects"
    )

    if not isinstance(
        objects,
        list,
    ):
        raise ValueError(
            "JSON file does not contain "
            "a valid 'objects' list."
        )

    # ========================================================
    # NO HARDCODED COUNT CHECK.
    # ========================================================

    if not objects:
        raise ValueError(
            "JSON file contains "
            "zero metadata objects."
        )

    ids = [
        obj.get("custom_id")
        for obj in objects
    ]

    if None in ids:
        raise ValueError(
            "At least one JSON object "
            "has no custom_id."
        )

    if len(ids) != len(
        set(ids)
    ):
        raise ValueError(
            "Duplicate custom_id "
            "found in JSON file."
        )

    print(
        f"[OK] JSON file validated: "
        f"{len(objects)} objects"
    )

    # --------------------------------------------------------
    # Connect
    # --------------------------------------------------------

    print(
        f"[INFO] Connecting to "
        f"Weaviate: {WEAVIATE_URL}"
    )

    client = (
        connect_to_weaviate()
    )

    # --------------------------------------------------------
    # Docker / Weaviate ready?
    # --------------------------------------------------------

    if not client.is_ready():

        raise RuntimeError(
            f"Weaviate is not ready "
            f"at {WEAVIATE_URL}.\n"
            f"Check your Docker "
            f"Weaviate container."
        )

    print(
        "[OK] Weaviate is ready."
    )

    # --------------------------------------------------------
    # custom_id property
    # --------------------------------------------------------

    ensure_custom_id_property(
        client
    )

    # --------------------------------------------------------
    # Insert / update
    # --------------------------------------------------------

    inserted = 0
    updated = 0

    client.batch.configure(
        batch_size=BATCH_SIZE
    )

    print()
    print(
        "[INFO] Loading objects "
        "into Weaviate..."
    )

    with client.batch as batch:

        for index, obj in enumerate(
            objects,
            start=1,
        ):

            custom_id = obj[
                "custom_id"
            ]

            properties = obj[
                "properties"
            ]

            # ------------------------------------------------
            # custom_id is NOT a UUID.
            #
            # Convert business ID into deterministic UUID.
            #
            # Same custom_id -> same UUID every run.
            # ------------------------------------------------

            object_uuid = str(

                uuid.uuid5(

                    uuid.NAMESPACE_URL,

                    (
                        f"{CLASS_NAME}:"
                        f"{custom_id}"
                    ),
                )
            )

            # ------------------------------------------------
            # Existing object -> replace/update
            # ------------------------------------------------

            if client.data_object.exists(

                object_uuid,

                class_name=CLASS_NAME,

            ):

                client.data_object.replace(

                    data_object=properties,

                    class_name=CLASS_NAME,

                    uuid=object_uuid,
                )

                updated += 1

            # ------------------------------------------------
            # New object -> batch insert
            # ------------------------------------------------

            else:

                batch.add_data_object(

                    data_object=properties,

                    class_name=CLASS_NAME,

                    uuid=object_uuid,
                )

                inserted += 1

            # ------------------------------------------------
            # Progress
            # ------------------------------------------------

            if (
                index % 50 == 0
                or index == len(objects)
            ):

                print(
                    f"[WEAVIATE] "
                    f"{index}/"
                    f"{len(objects)} "
                    f"processed"
                )

    # --------------------------------------------------------
    # Done
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print(
        "[OK] WEAVIATE LOAD COMPLETE"
    )
    print("=" * 70)

    print(
        f"Total objects : "
        f"{len(objects)}"
    )

    print(
        f"Inserted      : "
        f"{inserted}"
    )

    print(
        f"Updated       : "
        f"{updated}"
    )


# ============================================================
# SEARCH BY CUSTOM ID
# ============================================================

def search_by_custom_id(
    custom_id: str,
) -> Dict[str, Any]:

    client = (
        connect_to_weaviate()
    )

    if not client.is_ready():

        raise RuntimeError(
            "Weaviate is not ready."
        )

    result = (

        client.query

        .get(
            CLASS_NAME,
            [
                "custom_id",
                "columnName",
                "tableName",
                "module",
                "dataType",
                "description",
                "roleAccess",
                "used",
                "filterable",
            ],
        )

        .with_where(
            {
                "path": [
                    "custom_id"
                ],

                "operator":
                    "Equal",

                "valueText":
                    custom_id,
            }
        )

        .with_limit(5)

        .do()
    )

    return result


# ============================================================
# OPTIONAL SEARCH BY REPORT
# ============================================================

def search_by_module(
    module_name: str,
    limit: int = 100,
) -> Dict[str, Any]:

    client = (
        connect_to_weaviate()
    )

    if not client.is_ready():
        raise RuntimeError(
            "Weaviate is not ready."
        )

    result = (

        client.query

        .get(
            CLASS_NAME,
            [
                "custom_id",
                "columnName",
                "tableName",
                "module",
                "dataType",
                "description",
                "roleAccess",
            ],
        )

        .with_where(
            {
                "path": [
                    "module"
                ],

                "operator":
                    "Equal",

                "valueText":
                    module_name,
            }
        )

        .with_limit(limit)

        .do()
    )

    return result


# ============================================================
# MAIN
# ============================================================

def main() -> None:

    print("=" * 70)
    print(
        "STEP 1: "
        "EXCEL -> VALIDATED JSON"
    )
    print("=" * 70)

    excel_to_json(

        excel_path=(
            EXCEL_FILE
        ),

        json_path=(
            JSON_FILE
        ),
    )

    print()
    print("=" * 70)
    print(
        "STEP 2: "
        "JSON -> DOCKER WEAVIATE"
    )
    print("=" * 70)

    load_json_to_weaviate(
        JSON_FILE
    )

    # ========================================================
    # OPTIONAL TEST
    #
    # Uncomment after loading:
    # ========================================================

    # result = search_by_custom_id(
    #     "01.001.021"
    # )
    #
    # print(
    #     json.dumps(
    #         result,
    #         indent=2,
    #         ensure_ascii=False,
    #     )
    # )

    # ========================================================
    # OPTIONAL REPORT TEST
    # ========================================================

    # result = search_by_module(
    #     "Inventory Report"
    # )
    #
    # print(
    #     json.dumps(
    #         result,
    #         indent=2,
    #         ensure_ascii=False,
    #     )
    # )


if __name__ == "__main__":
    main()
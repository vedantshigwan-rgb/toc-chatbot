from __future__ import annotations

import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from sqlglot import exp, parse_one
from sqlglot.errors import ParseError

from services.sql_template_placeholders import SQLTemplatePlaceholderGuard

from services.extra_sql_filter_service import (
    SAFE_IDENTIFIER_PATTERN,
    _base_table_name,
    _direct_tables_for_select,
    _get_mapping_candidates,
    _normalize_key,
    _table_qualifier,
)


# ============================================================================
# CONFIGURATION
# ============================================================================

_BACKEND_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_ROLES_OPTIONS_WORKBOOK_CANDIDATES = (
    _BACKEND_ROOT / "Roles and options.xlsx",
    _BACKEND_ROOT / "Roles and options-F.xlsx",
)

DEFAULT_VARIABLE_MAPPING_WORKBOOK_CANDIDATES = (
    _BACKEND_ROOT / "redis_mapping.xlsx",
    _BACKEND_ROOT / "redis mapping.xlsx",
)

REQUIRED_ROLES_HEADERS = {
    "report name",
    "table name",
    "variable name",
}

PLACEHOLDER_NAME_PATTERN = re.compile(
    r"^[A-Za-z_][A-Za-z0-9_]*$"
)


# ============================================================================
# BASIC HELPERS
# ============================================================================


def _normalize_text(value: Any) -> str:
    return re.sub(
        r"\s+",
        " ",
        str(value or "").strip(),
    )


def _normalized_header(value: Any) -> str:
    return _normalize_text(value).casefold()


def _is_present_variable_value(value: Any) -> bool:
    if value is None:
        return False

    if isinstance(value, str):
        return bool(value.strip())

    if isinstance(value, (list, tuple, set)):
        return any(
            _is_present_variable_value(item)
            for item in value
        )

    return True


def _first_existing_path(
    candidates: tuple[Path, ...],
) -> Path:
    for path in candidates:
        if path.exists():
            return path.resolve()

    # Return the preferred path even when it does not exist so the caller
    # receives a useful FileNotFoundError containing the expected location.
    return candidates[0].resolve()


def _resolve_roles_workbook_path(
    roles_workbook_path: Optional[str],
) -> Path:
    configured = (
        roles_workbook_path
        or os.getenv("ROLES_OPTIONS_WORKBOOK")
    )

    if configured:
        return Path(configured).expanduser().resolve()

    return _first_existing_path(
        DEFAULT_ROLES_OPTIONS_WORKBOOK_CANDIDATES
    )


def _resolve_mapping_workbook_path(
    mapping_workbook_path: Optional[str],
) -> Path:
    configured = (
        mapping_workbook_path
        or os.getenv("VARIABLE_MAPPING_WORKBOOK")
    )

    if configured:
        return Path(configured).expanduser().resolve()

    return _first_existing_path(
        DEFAULT_VARIABLE_MAPPING_WORKBOOK_CANDIDATES
    )


# ============================================================================
# ROLES + OPTIONS WORKBOOK
# ============================================================================


@lru_cache(maxsize=8)
def _load_report_option_index(
    workbook_path: str,
    workbook_mtime_ns: int,
) -> Dict[str, tuple[Dict[str, Any], ...]]:
    """
    Read Roles and options.xlsx and build:

        normalized report name -> option variable definitions

    The sheet uses continuation rows for reports such as Docket Report, where
    the Report name cell is blank on the PERIOD TYPE / PERIOD NAME rows.
    Therefore the report name is intentionally forward-filled.

    Duplicate report/variable rows are merged. This matters for reports such
    as Inventory Report and Order Report which can appear under multiple
    navigation/sub-tab rows but describe the same report-level option.
    """

    workbook = load_workbook(
        workbook_path,
        data_only=True,
        read_only=True,
    )

    try:
        if not workbook.sheetnames:
            raise ValueError(
                "Roles and options workbook has no worksheets."
            )

        worksheet = workbook[workbook.sheetnames[0]]

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

        headers = [
            _normalized_header(value)
            for value in first_row
        ]

        header_indexes = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        missing_headers = sorted(
            REQUIRED_ROLES_HEADERS
            - set(header_indexes)
        )

        if missing_headers:
            raise ValueError(
                "Roles and options workbook is missing required headers: "
                + ", ".join(missing_headers)
            )

        option_header_index = header_indexes.get("options")

        by_report: Dict[
            str,
            Dict[str, Dict[str, Any]],
        ] = {}

        current_report_name = ""
        current_table_name = ""

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            raw_report_name = _normalize_text(
                row[header_indexes["report name"]]
            )
            raw_table_name = _normalize_text(
                row[header_indexes["table name"]]
            )
            variable_name = _normalize_text(
                row[header_indexes["variable name"]]
            )

            option_label = ""
            if option_header_index is not None:
                option_label = _normalize_text(
                    row[option_header_index]
                )

            if raw_report_name:
                current_report_name = raw_report_name

            if raw_table_name:
                current_table_name = raw_table_name

            if not current_report_name or not variable_name:
                continue

            report_key = _normalize_key(
                current_report_name
            )
            variable_key = _normalize_key(
                variable_name
            )

            report_bucket = by_report.setdefault(
                report_key,
                {},
            )

            existing = report_bucket.get(variable_key)

            if existing is None:
                existing = {
                    "report_name": current_report_name,
                    "variable_name": variable_name,
                    "option_labels": [],
                    "table_names": [],
                }
                report_bucket[variable_key] = existing

            if (
                option_label
                and option_label not in existing["option_labels"]
            ):
                existing["option_labels"].append(option_label)

            if (
                current_table_name
                and current_table_name not in existing["table_names"]
            ):
                existing["table_names"].append(current_table_name)

        return {
            report_key: tuple(variable_map.values())
            for report_key, variable_map in by_report.items()
        }

    finally:
        workbook.close()


def _get_report_option_definitions(
    *,
    report_name: str,
    roles_workbook_path: Optional[str] = None,
) -> tuple[Path, List[Dict[str, Any]]]:
    workbook_path = _resolve_roles_workbook_path(
        roles_workbook_path
    )

    if not workbook_path.exists():
        raise FileNotFoundError(
            "Roles and options workbook was not found at: "
            f"{workbook_path}"
        )

    mtime_ns = workbook_path.stat().st_mtime_ns
    index = _load_report_option_index(
        str(workbook_path),
        mtime_ns,
    )

    definitions = list(
        index.get(
            _normalize_key(report_name),
            (),
        )
    )

    return workbook_path, definitions


# ============================================================================
# SQL AST HELPERS
# ============================================================================


def _placeholder_pattern(variable_name: str) -> re.Pattern[str]:
    return re.compile(
        r"\{\s*"
        + re.escape(variable_name)
        + r"\s*\}",
        re.IGNORECASE,
    )


def _sql_contains_placeholder(
    sql_text: str,
    variable_name: str,
) -> bool:
    return bool(
        _placeholder_pattern(variable_name).search(
            str(sql_text or "")
        )
    )


def _build_contains_filter_condition(
    *,
    qualifier: str,
    column_name: str,
    variable_name: str,
    dialect: str,
) -> exp.Expression:
    if not SAFE_IDENTIFIER_PATTERN.fullmatch(column_name):
        raise ValueError(
            f"Unsafe mapped column name: {column_name}"
        )

    if (
        qualifier
        and not SAFE_IDENTIFIER_PATTERN.fullmatch(qualifier)
    ):
        raise ValueError(
            f"Unsafe SQL table alias/name: {qualifier}"
        )

    if not PLACEHOLDER_NAME_PATTERN.fullmatch(variable_name):
        raise ValueError(
            f"Unsafe variable name: {variable_name}"
        )

    left_side = (
        f"{qualifier}.{column_name}"
        if qualifier
        else column_name
    )

    # Actual values are intentionally not concatenated here. The existing
    # SQL injection stage replaces {variable_name} later.
    condition_sql = (
        f"{left_side} LIKE '%{{{variable_name}}}%'"
    )

    return parse_one(
        condition_sql,
        read=dialect,
    )


def _matches_mapped_column(
    column: exp.Column,
    *,
    select: exp.Select,
    column_name: str,
    qualifier_names: set[str],
) -> bool:
    # Do not let an outer SELECT rewrite predicates inside a nested SELECT.
    if column.find_ancestor(exp.Select) is not select:
        return False

    if _normalize_key(column.name) != _normalize_key(column_name):
        return False

    column_table = _normalize_key(column.table)

    # An unqualified mapped column is accepted because this SELECT scope has
    # already been proven to directly read the mapped physical table.
    if not column_table:
        return True

    return column_table in qualifier_names


def _expression_references_mapped_column(
    expression: exp.Expression,
    *,
    select: exp.Select,
    column_name: str,
    qualifier_names: set[str],
) -> bool:
    for column in expression.find_all(exp.Column):
        if _matches_mapped_column(
            column,
            select=select,
            column_name=column_name,
            qualifier_names=qualifier_names,
        ):
            return True

    return False


def _expression_contains_variable_placeholder(
    expression: exp.Expression,
    *,
    variable_name: str,
    dialect: str,
    placeholder_guard: SQLTemplatePlaceholderGuard,
) -> bool:
    try:
        sql_text = expression.sql(dialect=dialect)
    except Exception:
        sql_text = str(expression)

    return placeholder_guard.contains_variable(
        sql_text,
        variable_name,
    )


def _strip_target_predicates(
    expression: Optional[exp.Expression],
    *,
    select: exp.Select,
    column_name: str,
    qualifier_names: set[str],
    variable_name: str,
    dialect: str,
    require_placeholder: bool,
    placeholder_guard: SQLTemplatePlaceholderGuard,
) -> tuple[Optional[exp.Expression], int]:
    """
    Remove matching WHERE predicates while preserving the rest of the tree.

    Supplied value:
        require_placeholder=False
        -> remove every existing predicate on the mapped option column so the
           deterministic user-selected filter becomes authoritative.

    No supplied value:
        require_placeholder=True
        -> remove only a mapped-column predicate that actually contains the
           exact {variable_name} placeholder. Hardcoded conditions such as
           ORG_NAME = 'ACD-SPARES' are intentionally preserved.
    """

    if expression is None:
        return None, 0

    if isinstance(expression, exp.Paren):
        cleaned_inner, removed = _strip_target_predicates(
            expression.this,
            select=select,
            column_name=column_name,
            qualifier_names=qualifier_names,
            variable_name=variable_name,
            dialect=dialect,
            require_placeholder=require_placeholder,
            placeholder_guard=placeholder_guard,
        )

        if cleaned_inner is None:
            return None, removed

        return exp.Paren(this=cleaned_inner), removed

    if isinstance(expression, (exp.And, exp.Or)):
        cleaned_left, removed_left = _strip_target_predicates(
            expression.this,
            select=select,
            column_name=column_name,
            qualifier_names=qualifier_names,
            variable_name=variable_name,
            dialect=dialect,
            require_placeholder=require_placeholder,
            placeholder_guard=placeholder_guard,
        )

        cleaned_right, removed_right = _strip_target_predicates(
            expression.expression,
            select=select,
            column_name=column_name,
            qualifier_names=qualifier_names,
            variable_name=variable_name,
            dialect=dialect,
            require_placeholder=require_placeholder,
            placeholder_guard=placeholder_guard,
        )

        removed = removed_left + removed_right

        if cleaned_left is None and cleaned_right is None:
            return None, removed

        if cleaned_left is None:
            return cleaned_right, removed

        if cleaned_right is None:
            return cleaned_left, removed

        if isinstance(expression, exp.And):
            return exp.and_(cleaned_left, cleaned_right), removed

        return exp.or_(cleaned_left, cleaned_right), removed

    references_column = _expression_references_mapped_column(
        expression,
        select=select,
        column_name=column_name,
        qualifier_names=qualifier_names,
    )

    if not references_column:
        return expression.copy(), 0

    if require_placeholder:
        if not _expression_contains_variable_placeholder(
            expression,
            variable_name=variable_name,
            dialect=dialect,
            placeholder_guard=placeholder_guard,
        ):
            return expression.copy(), 0

    return None, 1


# ============================================================================
# VARIABLE / MAPPING RESOLUTION
# ============================================================================


def _request_variable_entries(
    variables: Optional[List[Dict[str, Any]]],
    variable_name: str,
) -> List[Dict[str, Any]]:
    target_key = _normalize_key(variable_name)
    matches: List[Dict[str, Any]] = []

    for item in variables or []:
        if not isinstance(item, dict):
            continue

        if _normalize_key(item.get("name")) != target_key:
            continue

        matches.append(dict(item))

    return matches


def _has_supplied_value(entries: List[Dict[str, Any]]) -> bool:
    return any(
        _is_present_variable_value(item.get("value"))
        for item in entries
    )


def _candidate_matches_tables(
    candidate: Dict[str, Any],
    table_names: set[str],
) -> bool:
    if not table_names:
        return True

    candidate_table = _normalize_key(
        _base_table_name(candidate.get("table_name"))
    )

    return candidate_table in table_names


def _matching_selects_for_candidate(
    expression: exp.Expression,
    candidate: Dict[str, Any],
) -> List[tuple[exp.Select, List[exp.Table]]]:
    mapped_base_table = _normalize_key(
        _base_table_name(candidate.get("table_name"))
    )

    matching_selects: List[
        tuple[exp.Select, List[exp.Table]]
    ] = []

    for select in expression.find_all(exp.Select):
        direct_tables = _direct_tables_for_select(select)

        matching_tables = [
            table
            for table in direct_tables
            if _normalize_key(
                _base_table_name(table.name)
            ) == mapped_base_table
        ]

        if matching_tables:
            matching_selects.append(
                (select, matching_tables)
            )

    return matching_selects


# ============================================================================
# PUBLIC SERVICE
# ============================================================================


def apply_report_option_sql_filters(
    *,
    sql_template: str,
    report_name: str,
    variables: List[Dict[str, Any]],
    allowed_tables: Optional[List[str]] = None,
    roles_workbook_path: Optional[str] = None,
    mapping_workbook_path: Optional[str] = None,
    dialect: str = "oracle",
) -> Dict[str, Any]:
    """
    Apply report-option rules after LLM SQL generation.

    Source of option variables:
        Roles and options.xlsx
        selected report -> configured variable names

    Source of physical SQL mappings:
        redis_mapping.xlsx / redis mapping.xlsx
        report + variable -> mapped table + mapped column

    Rules for every configured option variable:

    1. If the request supplies a non-empty value (source=input, option,
       feedback, or any other source):
         - remove any existing predicate on the mapped column
         - add one authoritative predicate:
               MAPPED_COLUMN LIKE '%{variable_name}%'

    2. If the request does NOT supply a value:
         - remove only mapped-column predicates that contain that variable's
           placeholder, e.g. ORG_NAME = '{org_name}' or
           ORG_NAME LIKE '%{org_name}%'
         - preserve hardcoded predicates, e.g. ORG_NAME = 'ACD-SPARES'

    The function never injects literal user values. The normal SQL injection
    stage runs afterwards.
    """

    cleaned_sql = str(sql_template or "").strip()

    if not cleaned_sql:
        return {
            "success": False,
            "message": "SQL template is empty.",
            "sql_template": cleaned_sql,
            "applied": False,
        }

    try:
        roles_path, option_definitions = _get_report_option_definitions(
            report_name=report_name,
            roles_workbook_path=roles_workbook_path,
        )
    except Exception as error:
        return {
            "success": False,
            "message": (
                "Could not load report option variables from "
                "Roles and options.xlsx."
            ),
            "sql_template": cleaned_sql,
            "applied": False,
            "exception_type": type(error).__name__,
            "details": {
                "internal_message": str(error),
            },
        }

    configured_variable_names = [
        str(item.get("variable_name") or "").strip()
        for item in option_definitions
        if str(item.get("variable_name") or "").strip()
    ]

    if not option_definitions:
        return {
            "success": True,
            "message": (
                f"No report option variables are configured for '{report_name}'."
            ),
            "sql_template": cleaned_sql,
            "applied": False,
            "configured_variable_names": [],
            "roles_workbook_path": str(roles_path),
            "mapping_workbook_path": None,
            "variable_results": [],
        }

    mapping_path = _resolve_mapping_workbook_path(
        mapping_workbook_path
    )

    if not mapping_path.exists():
        return {
            "success": False,
            "message": (
                "Variable mapping workbook was not found at: "
                f"{mapping_path}"
            ),
            "sql_template": cleaned_sql,
            "applied": False,
            "configured_variable_names": configured_variable_names,
            "roles_workbook_path": str(roles_path),
            "mapping_workbook_path": str(mapping_path),
        }

    had_semicolon = cleaned_sql.endswith(";")
    sql_without_semicolon = cleaned_sql.rstrip(";").strip()

    # Protect numeric/unquoted templates such as > {performance_index}.
    # Without this, SQLGlot can emit > STRUCT(performance_index).
    placeholder_guard = SQLTemplatePlaceholderGuard(sql_without_semicolon)

    try:
        expression = parse_one(
            placeholder_guard.protected_sql,
            read=dialect,
        )
    except ParseError as error:
        return {
            "success": False,
            "message": (
                "Generated SQL could not be parsed before applying "
                "report option filters."
            ),
            "sql_template": cleaned_sql,
            "applied": False,
            "exception_type": type(error).__name__,
            "details": {
                "parse_error": str(error),
            },
        }

    allowed_base_tables = {
        _normalize_key(_base_table_name(table_name))
        for table_name in allowed_tables or []
        if _base_table_name(table_name)
    }

    variable_results: List[Dict[str, Any]] = []
    any_change = False

    for option_definition in option_definitions:
        variable_name = str(
            option_definition.get("variable_name") or ""
        ).strip()

        if not variable_name:
            continue

        request_entries = _request_variable_entries(
            variables,
            variable_name,
        )
        value_supplied = _has_supplied_value(
            request_entries
        )

        # For an absent variable, there is nothing to clean if the LLM did
        # not generate that placeholder at all. This also prevents an
        # irrelevant missing mapping row from breaking otherwise valid SQL.
        placeholder_present_in_sql = _sql_contains_placeholder(
            sql_without_semicolon,
            variable_name,
        )

        if not value_supplied and not placeholder_present_in_sql:
            variable_results.append(
                {
                    "variable_name": variable_name,
                    "option_labels": list(
                        option_definition.get("option_labels") or []
                    ),
                    "roles_table_names": list(
                        option_definition.get("table_names") or []
                    ),
                    "value_supplied": False,
                    "action": "no_value_no_placeholder_no_change",
                    "changed": False,
                    "removed_predicate_count": 0,
                }
            )
            continue

        try:
            _, candidates = _get_mapping_candidates(
                report_name=report_name,
                variable_name=variable_name,
                mapping_workbook_path=str(mapping_path),
            )
        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Could not load the report-wise SQL mapping for "
                    f"option variable '{variable_name}'."
                ),
                "sql_template": cleaned_sql,
                "applied": any_change,
                "configured_variable_names": configured_variable_names,
                "roles_workbook_path": str(roles_path),
                "mapping_workbook_path": str(mapping_path),
                "variable_results": variable_results,
                "exception_type": type(error).__name__,
                "details": {
                    "internal_message": str(error),
                },
            }

        roles_table_names = {
            _normalize_key(_base_table_name(table_name))
            for table_name in (
                option_definition.get("table_names") or []
            )
            if _base_table_name(table_name)
        }

        if roles_table_names:
            candidates = [
                candidate
                for candidate in candidates
                if _candidate_matches_tables(
                    candidate,
                    roles_table_names,
                )
            ]

        if allowed_base_tables:
            candidates = [
                candidate
                for candidate in candidates
                if _candidate_matches_tables(
                    candidate,
                    allowed_base_tables,
                )
            ]

        if not candidates:
            return {
                "success": False,
                "message": (
                    f"No trusted mapping exists for option variable "
                    f"'{variable_name}' in report '{report_name}'."
                ),
                "sql_template": cleaned_sql,
                "applied": any_change,
                "configured_variable_names": configured_variable_names,
                "roles_workbook_path": str(roles_path),
                "mapping_workbook_path": str(mapping_path),
                "variable_results": variable_results,
                "details": {
                    "roles_table_names": sorted(roles_table_names),
                    "allowed_tables": allowed_tables or [],
                },
            }

        matching_candidates: List[
            tuple[
                Dict[str, Any],
                List[tuple[exp.Select, List[exp.Table]]],
            ]
        ] = []

        for candidate in candidates:
            matching_selects = _matching_selects_for_candidate(
                expression,
                candidate,
            )

            if matching_selects:
                matching_candidates.append(
                    (candidate, matching_selects)
                )

        if not matching_candidates:
            return {
                "success": False,
                "message": (
                    "Generated SQL does not use the table mapped to "
                    f"option variable '{variable_name}' for report "
                    f"'{report_name}'."
                ),
                "sql_template": cleaned_sql,
                "applied": any_change,
                "configured_variable_names": configured_variable_names,
                "roles_workbook_path": str(roles_path),
                "mapping_workbook_path": str(mapping_path),
                "variable_results": variable_results,
                "details": {
                    "mapping_candidates": candidates,
                },
            }

        # Multiple identical mappings can exist in Excel. Treat them as one;
        # genuinely different mapped table/column pairs remain ambiguous.
        unique_matches: Dict[
            tuple[str, str],
            tuple[
                Dict[str, Any],
                List[tuple[exp.Select, List[exp.Table]]],
            ],
        ] = {}

        for candidate, matching_selects in matching_candidates:
            key = (
                _normalize_key(
                    _base_table_name(candidate.get("table_name"))
                ),
                _normalize_key(candidate.get("column_name")),
            )
            unique_matches[key] = (
                candidate,
                matching_selects,
            )

        if len(unique_matches) != 1:
            return {
                "success": False,
                "message": (
                    "Multiple physical mappings match option variable "
                    f"'{variable_name}' in the generated SQL."
                ),
                "sql_template": cleaned_sql,
                "applied": any_change,
                "configured_variable_names": configured_variable_names,
                "roles_workbook_path": str(roles_path),
                "mapping_workbook_path": str(mapping_path),
                "variable_results": variable_results,
                "details": {
                    "matching_mappings": [
                        candidate
                        for candidate, _ in matching_candidates
                    ],
                },
            }

        selected_mapping, matching_selects = next(
            iter(unique_matches.values())
        )

        column_name = str(
            selected_mapping.get("column_name") or ""
        ).strip()

        removed_predicate_count = 0
        updated_select_count = 0

        for select, matching_tables in matching_selects:
            if len(matching_tables) != 1:
                return {
                    "success": False,
                    "message": (
                        f"Mapped table for option variable '{variable_name}' "
                        "is used more than once in the same SELECT scope, "
                        "so the correct alias is ambiguous."
                    ),
                    "sql_template": cleaned_sql,
                    "applied": any_change,
                    "configured_variable_names": configured_variable_names,
                    "roles_workbook_path": str(roles_path),
                    "mapping_workbook_path": str(mapping_path),
                    "variable_results": variable_results,
                }

            table = matching_tables[0]
            qualifier = _table_qualifier(table)

            qualifier_names = {
                _normalize_key(qualifier),
                _normalize_key(table.name),
                _normalize_key(_base_table_name(table.name)),
            }
            qualifier_names.discard("")

            existing_where = select.args.get("where")

            if existing_where is None:
                base_condition = None
                removed_here = 0
            else:
                base_condition, removed_here = _strip_target_predicates(
                    existing_where.this,
                    select=select,
                    column_name=column_name,
                    qualifier_names=qualifier_names,
                    variable_name=variable_name,
                    dialect=dialect,
                    require_placeholder=(not value_supplied),
                    placeholder_guard=placeholder_guard,
                )

            removed_predicate_count += removed_here

            if value_supplied:
                try:
                    enforced_condition = _build_contains_filter_condition(
                        qualifier=qualifier,
                        column_name=column_name,
                        variable_name=variable_name,
                        dialect=dialect,
                    )
                except Exception as error:
                    return {
                        "success": False,
                        "message": (
                            "Could not build SQL filter for option variable "
                            f"'{variable_name}'."
                        ),
                        "sql_template": cleaned_sql,
                        "applied": any_change,
                        "configured_variable_names": configured_variable_names,
                        "roles_workbook_path": str(roles_path),
                        "mapping_workbook_path": str(mapping_path),
                        "variable_results": variable_results,
                        "exception_type": type(error).__name__,
                        "details": {
                            "internal_message": str(error),
                        },
                    }

                if base_condition is None:
                    new_where_condition = enforced_condition
                else:
                    new_where_condition = exp.and_(
                        base_condition,
                        enforced_condition,
                    )

                select.set(
                    "where",
                    exp.Where(this=new_where_condition),
                )
                updated_select_count += 1
                any_change = True

            elif removed_here:
                if base_condition is None:
                    select.set("where", None)
                else:
                    select.set(
                        "where",
                        exp.Where(this=base_condition),
                    )

                updated_select_count += 1
                any_change = True

        if value_supplied:
            action = "enforced_supplied_value"
        elif removed_predicate_count:
            action = "removed_unbound_placeholder_predicate"
        else:
            action = "no_matching_placeholder_predicate"

        variable_results.append(
            {
                "variable_name": variable_name,
                "option_labels": list(
                    option_definition.get("option_labels") or []
                ),
                "roles_table_names": list(
                    option_definition.get("table_names") or []
                ),
                "value_supplied": value_supplied,
                "variable_sources": sorted(
                    {
                        str(item.get("source") or "input")
                        .strip()
                        .lower()
                        for item in request_entries
                    }
                ),
                "action": action,
                "changed": bool(
                    value_supplied or removed_predicate_count
                ),
                "mapped_column_name": column_name,
                "mapped_table_name": str(
                    selected_mapping.get("table_name") or ""
                ).strip(),
                "storage_type": str(
                    selected_mapping.get("storage_type") or ""
                ).strip(),
                "removed_predicate_count": removed_predicate_count,
                "select_scopes_updated": updated_select_count,
            }
        )

    # Restore exact placeholder spelling before variable injection.
    try:
        final_sql_template = placeholder_guard.restore(
            expression.sql(dialect=dialect).strip()
        )
    except ValueError as error:
        return {
            "success": False,
            "message": "SQL template placeholders could not be restored.",
            "sql_template": cleaned_sql,
            "applied": False,
            "exception_type": type(error).__name__,
        }

    if had_semicolon:
        final_sql_template += ";"

    return {
        "success": True,
        "message": (
            "Report option SQL rules were applied successfully."
        ),
        "sql_template": final_sql_template,
        "applied": any_change,
        "configured_variable_names": configured_variable_names,
        "roles_workbook_path": str(roles_path),
        "mapping_workbook_path": str(mapping_path),
        "variable_results": variable_results,
    }

from __future__ import annotations

import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from sqlglot import exp, parse_one
from sqlglot.errors import ParseError


# ============================================================================
# CONFIGURATION
# ============================================================================

DEFAULT_VARIABLE_MAPPING_WORKBOOK = (
    Path(__file__).resolve().parents[1]
    / "redis mapping.xlsx"
)

REQUIRED_MAPPING_HEADERS = {
    "report name",
    "variables",
    "column names",
    "table name",
    "storage_type",
}

SAFE_IDENTIFIER_PATTERN = re.compile(
    r"^[A-Za-z_][A-Za-z0-9_$#]*$"
)


# ============================================================================
# MAPPING HELPERS
# ============================================================================

def _normalize_text(
    value: Any,
) -> str:
    return re.sub(
        r"\s+",
        " ",
        str(value or "").strip(),
    )


def _normalize_key(
    value: Any,
) -> str:
    return (
        _normalize_text(value)
        .casefold()
    )


def _base_table_name(
    value: Any,
) -> str:
    """
    KPCL_REPORTING.FULLKIT_ITEMWISE_INV
        -> FULLKIT_ITEMWISE_INV
    """

    return (
        str(value or "")
        .strip()
        .split(".")[-1]
        .strip()
        .strip('"')
    )


def _resolve_mapping_workbook_path(
    mapping_workbook_path: Optional[str],
) -> Path:
    configured_path = (
        mapping_workbook_path
        or os.getenv(
            "VARIABLE_MAPPING_WORKBOOK"
        )
    )

    if configured_path:
        return (
            Path(configured_path)
            .expanduser()
            .resolve()
        )

    return (
        DEFAULT_VARIABLE_MAPPING_WORKBOOK
        .resolve()
    )


@lru_cache(maxsize=8)
def _load_mapping_index(
    workbook_path: str,
    workbook_mtime_ns: int,
) -> Dict[
    tuple[str, str],
    tuple[Dict[str, str], ...],
]:
    """
    Build:

        (report_name, variable_name)
            -> mapping rows

    The file modification time is part of the cache key, so
    editing the workbook causes it to be reloaded automatically.
    """

    workbook = load_workbook(
        workbook_path,
        data_only=True,
        read_only=True,
    )

    try:
        if not workbook.sheetnames:
            raise ValueError(
                "The variable mapping workbook "
                "has no worksheets."
            )

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

        headers = [
            _normalize_key(value)
            for value in first_row
        ]

        header_indexes = {
            header: index
            for index, header
            in enumerate(headers)
            if header
        }

        missing_headers = sorted(
            REQUIRED_MAPPING_HEADERS
            - set(header_indexes)
        )

        if missing_headers:
            raise ValueError(
                "The variable mapping workbook "
                "is missing required headers: "
                + ", ".join(
                    missing_headers
                )
            )

        mapping_index: Dict[
            tuple[str, str],
            List[Dict[str, str]],
        ] = {}

        seen_rows = set()

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            report_name = _normalize_text(
                row[
                    header_indexes[
                        "report name"
                    ]
                ]
            )

            variable_name = _normalize_text(
                row[
                    header_indexes[
                        "variables"
                    ]
                ]
            )

            column_name = _normalize_text(
                row[
                    header_indexes[
                        "column names"
                    ]
                ]
            )

            table_name = _normalize_text(
                row[
                    header_indexes[
                        "table name"
                    ]
                ]
            )

            storage_type = _normalize_text(
                row[
                    header_indexes[
                        "storage_type"
                    ]
                ]
            )

            if not (
                report_name
                and variable_name
                and column_name
                and table_name
            ):
                continue

            key = (
                _normalize_key(
                    report_name
                ),
                _normalize_key(
                    variable_name
                ),
            )

            row_identity = (
                key,
                _normalize_key(
                    column_name
                ),
                _normalize_key(
                    table_name
                ),
                _normalize_key(
                    storage_type
                ),
            )

            if row_identity in seen_rows:
                continue

            seen_rows.add(
                row_identity
            )

            mapping_index.setdefault(
                key,
                [],
            ).append(
                {
                    "report_name":
                        report_name,

                    "variable_name":
                        variable_name,

                    "column_name":
                        column_name,

                    "table_name":
                        table_name,

                    "storage_type":
                        storage_type,
                }
            )

        return {
            key: tuple(rows)
            for key, rows
            in mapping_index.items()
        }

    finally:
        workbook.close()


def _get_mapping_candidates(
    *,
    report_name: str,
    variable_name: str,
    mapping_workbook_path: Optional[str],
) -> tuple[
    Path,
    List[Dict[str, str]],
]:
    workbook_path = (
        _resolve_mapping_workbook_path(
            mapping_workbook_path
        )
    )

    if not workbook_path.is_file():
        raise FileNotFoundError(
            "Variable mapping workbook "
            "not found: "
            f"{workbook_path}"
        )

    mapping_index = (
        _load_mapping_index(
            str(workbook_path),
            workbook_path
            .stat()
            .st_mtime_ns,
        )
    )

    candidates = list(
        mapping_index.get(
            (
                _normalize_key(
                    report_name
                ),
                _normalize_key(
                    variable_name
                ),
            ),
            (),
        )
    )

    return (
        workbook_path,
        candidates,
    )


# ============================================================================
# SQL HELPERS
# ============================================================================

def _direct_tables_for_select(
    select: exp.Select,
) -> List[exp.Table]:
    """
    Return tables that belong directly to this SELECT scope.

    Tables inside nested subqueries belong to their own SELECT
    and are therefore excluded here.
    """

    direct_tables: List[
        exp.Table
    ] = []

    for table in select.find_all(
        exp.Table
    ):
        nearest_select = (
            table.find_ancestor(
                exp.Select
            )
        )

        if nearest_select is select:
            direct_tables.append(
                table
            )

    return direct_tables


def _table_qualifier(
    table: exp.Table,
) -> str:
    """
    Prefer an SQL alias when one exists.

    Example:
        FULLKIT_ITEMWISE_INV f
            -> f

        FULLKIT_ITEMWISE_INV
            -> FULLKIT_ITEMWISE_INV
    """

    alias = str(
        table.alias or ""
    ).strip()

    if alias:
        return alias

    return str(
        table.name or ""
    ).strip()


def _build_filter_condition(
    *,
    qualifier: str,
    column_name: str,
    variable_name: str,
    dialect: str,
):
    """
    Build:

        TABLE_ALIAS.ORG_NAME = '{org_name}'

    The actual value is NOT inserted here. The existing
    sql_injection_service will insert it later.
    """

    if not SAFE_IDENTIFIER_PATTERN.fullmatch(
        column_name
    ):
        raise ValueError(
            "Unsafe mapped column name: "
            f"{column_name}"
        )

    if (
        qualifier
        and not SAFE_IDENTIFIER_PATTERN.fullmatch(
            qualifier
        )
    ):
        raise ValueError(
            "Unsafe SQL table alias/name: "
            f"{qualifier}"
        )

    if not SAFE_IDENTIFIER_PATTERN.fullmatch(
        variable_name
    ):
        raise ValueError(
            "Unsafe variable name: "
            f"{variable_name}"
        )

    left_side = (
        f"{qualifier}.{column_name}"
        if qualifier
        else column_name
    )

    condition_sql = (
        f"{left_side} = "
        f"'{{{variable_name}}}'"
    )

    return parse_one(
        condition_sql,
        read=dialect,
    )


# ============================================================================
# PUBLIC SERVICE
# ============================================================================

def apply_extra_sql_filters(
    *,
    sql_template: str,
    report_name: str,
    extra_filters: List[
        Dict[str, Any]
    ],
    allowed_tables: Optional[
        List[str]
    ] = None,
    mapping_workbook_path: Optional[
        str
    ] = None,
    dialect: str = "oracle",
) -> Dict[str, Any]:
    """
    Add sidebar option variables to an already-generated SQL template.

    IMPORTANT:
    - Extra filters do NOT participate in FAISS.
    - Extra filters do NOT participate in base SQL generation.
    - This service only runs AFTER base SQL generation.
    - It resolves report + variable -> table + column from the
      backend mapping workbook.
    - If WHERE exists, the new condition is added with AND.
    - If WHERE does not exist, a new WHERE is created.
    - GROUP BY / HAVING / ORDER BY remain in the correct position.

    Example:

        Base SQL:
            SELECT DISTINCT BUYER_NAME
            FROM KPCL_REPORTING.FULLKIT_ITEMWISE_INV
            ORDER BY BUYER_NAME;

        Extra filter:
            org_name = ACD - CENTRI

        Mapping:
            org_name
                -> FULLKIT_ITEMWISE_INV.ORG_NAME

        Final template:
            SELECT DISTINCT BUYER_NAME
            FROM KPCL_REPORTING.FULLKIT_ITEMWISE_INV
            WHERE FULLKIT_ITEMWISE_INV.ORG_NAME = '{org_name}'
            ORDER BY BUYER_NAME;
    """

    cleaned_sql = str(
        sql_template or ""
    ).strip()

    if not cleaned_sql:
        return {
            "success": False,
            "message": (
                "SQL template is empty."
            ),
            "sql_template": cleaned_sql,
            "applied_filters": [],
        }

    normalized_filters: List[
        Dict[str, Any]
    ] = []

    for item in extra_filters or []:
        if not isinstance(
            item,
            dict,
        ):
            continue

        variable_name = str(
            item.get("name")
            or ""
        ).strip().lower()

        variable_value = item.get(
            "value"
        )

        if (
            not variable_name
            or variable_value is None
        ):
            continue

        normalized_filters.append(
            {
                "name":
                    variable_name,

                "value":
                    variable_value,

                "source":
                    "option",
            }
        )

    # Nothing to add.
    if not normalized_filters:
        return {
            "success": True,
            "message": (
                "No extra SQL filters "
                "were supplied."
            ),
            "sql_template": cleaned_sql,
            "applied_filters": [],
            "mapping_workbook_path": None,
        }

    # ----------------------------------------------------------------
    # Parse the generated SQL once.
    # ----------------------------------------------------------------

    had_semicolon = (
        cleaned_sql.endswith(";")
    )

    sql_without_semicolon = (
        cleaned_sql
        .rstrip(";")
        .strip()
    )

    try:
        expression = parse_one(
            sql_without_semicolon,
            read=dialect,
        )

    except ParseError as error:
        return {
            "success": False,
            "message": (
                "Generated SQL could not "
                "be parsed before applying "
                "option filters."
            ),
            "sql_template": cleaned_sql,
            "applied_filters": [],
            "exception_type": (
                type(error).__name__
            ),
            "details": {
                "parse_error":
                    str(error),
            },
        }

    allowed_base_tables = {
        _normalize_key(
            _base_table_name(
                table_name
            )
        )
        for table_name
        in allowed_tables or []
        if _base_table_name(
            table_name
        )
    }

    applied_filters: List[
        Dict[str, Any]
    ] = []

    workbook_path_used = None

    # ----------------------------------------------------------------
    # Apply each sidebar option filter.
    # ----------------------------------------------------------------

    for extra_filter in (
        normalized_filters
    ):
        variable_name = (
            extra_filter["name"]
        )

        # ------------------------------------------------------------
        # 1. Resolve report + variable -> table + column.
        # ------------------------------------------------------------

        try:
            (
                workbook_path,
                candidates,
            ) = _get_mapping_candidates(
                report_name=(
                    report_name
                ),
                variable_name=(
                    variable_name
                ),
                mapping_workbook_path=(
                    mapping_workbook_path
                ),
            )

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Could not load SQL "
                    "mapping for option "
                    f"variable '{variable_name}'."
                ),
                "sql_template":
                    cleaned_sql,

                "applied_filters":
                    applied_filters,

                "exception_type":
                    type(error).__name__,

                "details": {
                    "internal_message":
                        str(error),
                },
            }

        workbook_path_used = str(
            workbook_path
        )

        if not candidates:
            return {
                "success": False,
                "message": (
                    "No SQL mapping exists "
                    f"for option variable "
                    f"'{variable_name}' in "
                    f"report '{report_name}'."
                ),
                "sql_template":
                    cleaned_sql,

                "applied_filters":
                    applied_filters,

                "mapping_workbook_path":
                    workbook_path_used,

                "details": {
                    "report_name":
                        report_name,

                    "variable_name":
                        variable_name,
                },
            }

        # ------------------------------------------------------------
        # 2. Mapping must point to one of the trusted FAQ tables.
        # ------------------------------------------------------------

        if allowed_base_tables:
            candidates = [
                candidate
                for candidate
                in candidates
                if (
                    _normalize_key(
                        _base_table_name(
                            candidate.get(
                                "table_name"
                            )
                        )
                    )
                    in allowed_base_tables
                )
            ]

        if not candidates:
            return {
                "success": False,
                "message": (
                    "The SQL mapping for "
                    f"option variable "
                    f"'{variable_name}' does "
                    "not target a trusted "
                    "table for this report."
                ),
                "sql_template":
                    cleaned_sql,

                "applied_filters":
                    applied_filters,

                "mapping_workbook_path":
                    workbook_path_used,

                "details": {
                    "allowed_tables":
                        allowed_tables or [],
                },
            }

        # ------------------------------------------------------------
        # 3. Candidate table must actually occur in generated SQL.
        # ------------------------------------------------------------

        matching_candidates = []

        for candidate in candidates:
            mapped_base_table = (
                _normalize_key(
                    _base_table_name(
                        candidate.get(
                            "table_name"
                        )
                    )
                )
            )

            matching_selects = []

            for select in expression.find_all(
                exp.Select
            ):
                direct_tables = (
                    _direct_tables_for_select(
                        select
                    )
                )

                matching_tables = [
                    table
                    for table
                    in direct_tables
                    if (
                        _normalize_key(
                            _base_table_name(
                                table.name
                            )
                        )
                        ==
                        mapped_base_table
                    )
                ]

                if matching_tables:
                    matching_selects.append(
                        (
                            select,
                            matching_tables,
                        )
                    )

            if matching_selects:
                matching_candidates.append(
                    (
                        candidate,
                        matching_selects,
                    )
                )

        if not matching_candidates:
            return {
                "success": False,
                "message": (
                    "The generated SQL does "
                    "not use the mapped table "
                    f"for option variable "
                    f"'{variable_name}'."
                ),
                "sql_template":
                    cleaned_sql,

                "applied_filters":
                    applied_filters,

                "mapping_workbook_path":
                    workbook_path_used,

                "details": {
                    "mapping_candidates":
                        candidates,
                },
            }

        if (
            len(matching_candidates)
            > 1
        ):
            return {
                "success": False,
                "message": (
                    "Multiple SQL mappings "
                    f"match option variable "
                    f"'{variable_name}'."
                ),
                "sql_template":
                    cleaned_sql,

                "applied_filters":
                    applied_filters,

                "mapping_workbook_path":
                    workbook_path_used,

                "details": {
                    "matching_mappings": [
                        candidate
                        for (
                            candidate,
                            _
                        )
                        in matching_candidates
                    ],
                },
            }

        (
            selected_mapping,
            matching_selects,
        ) = matching_candidates[0]

        column_name = str(
            selected_mapping.get(
                "column_name"
            )
            or ""
        ).strip()

        updated_select_count = 0

        # ------------------------------------------------------------
        # 4. Add WHERE or AND in every SELECT scope that directly
        #    reads the mapped physical table.
        # ------------------------------------------------------------

        for (
            select,
            matching_tables,
        ) in matching_selects:

            # Self-join of the same mapped table is ambiguous.
            if len(matching_tables) != 1:
                return {
                    "success": False,
                    "message": (
                        "The mapped table for "
                        f"option variable "
                        f"'{variable_name}' is "
                        "used more than once in "
                        "the same SELECT scope."
                    ),
                    "sql_template":
                        cleaned_sql,

                    "applied_filters":
                        applied_filters,

                    "mapping_workbook_path":
                        workbook_path_used,
                }

            table = matching_tables[0]

            qualifier = (
                _table_qualifier(
                    table
                )
            )

            try:
                condition = (
                    _build_filter_condition(
                        qualifier=qualifier,
                        column_name=(
                            column_name
                        ),
                        variable_name=(
                            variable_name
                        ),
                        dialect=dialect,
                    )
                )

            except Exception as error:
                return {
                    "success": False,
                    "message": (
                        "Could not build SQL "
                        f"filter for option "
                        f"variable "
                        f"'{variable_name}'."
                    ),
                    "sql_template":
                        cleaned_sql,

                    "applied_filters":
                        applied_filters,

                    "exception_type":
                        type(error).__name__,

                    "details": {
                        "internal_message":
                            str(error),
                    },
                }

            existing_where = (
                select.args.get(
                    "where"
                )
            )

            # No WHERE:
            #
            # FROM TABLE
            # WHERE ORG_NAME = '{org_name}'
            if existing_where is None:
                select.set(
                    "where",
                    exp.Where(
                        this=condition
                    ),
                )

            # Existing WHERE:
            #
            # WHERE OLD_CONDITION
            #   AND ORG_NAME = '{org_name}'
            else:
                select.set(
                    "where",
                    exp.Where(
                        this=exp.and_(
                            existing_where.this,
                            condition,
                        )
                    ),
                )

            updated_select_count += 1

        applied_filters.append(
            {
                "variable_name":
                    variable_name,

                "column_name":
                    column_name,

                "table_name":
                    str(
                        selected_mapping.get(
                            "table_name"
                        )
                        or ""
                    ).strip(),

                "storage_type":
                    str(
                        selected_mapping.get(
                            "storage_type"
                        )
                        or ""
                    ).strip(),

                "select_scopes_updated":
                    updated_select_count,
            }
        )

    # ----------------------------------------------------------------
    # Serialize back to Oracle SQL.
    # ----------------------------------------------------------------

    final_sql_template = (
        expression.sql(
            dialect=dialect
        ).strip()
    )

    if had_semicolon:
        final_sql_template += ";"

    return {
        "success": True,
        "message": (
            "Extra SQL filters were "
            "applied successfully."
        ),
        "sql_template":
            final_sql_template,

        "applied_filters":
            applied_filters,

        "mapping_workbook_path":
            workbook_path_used,
    }
from __future__ import annotations

import asyncio
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook


DEFAULT_COLUMNS = [
    "Candidate Key",
    "Created At UTC",
    "Feedback Type",
    "Thumb Rating",
    "Review Priority",
    "Report Name",
    "Original User Question",
    "Previous Normalized Question",
    "Feedback Input",
    "Normalized Feedback",
    "New Normalized Question",
    "Normalized Question",
    "Variable Names",
    "Variables Key",
    "Variables JSON",
    "Action",
    "Selected Example",
    "FAISS Examples Used",
    "SQL Template",
    "Executed SQL",
    "Query Executed",
    "Result Count",
    "Results JSON",
    "Assistant Answer",
    "Request Mode",
    "Chat ID",
    "Source User Message ID",
    "Source Assistant Message ID",
    "Workflow ID",
    "Diagnostics JSON",
    "Review Status",
    "Display Question",
    "Review Notes",
]


class FlaggedFaqService:
    """
    Best-effort Excel audit queue for FAQ-learning candidates.

    The workbook collects two kinds of signals:
      1. Modify/rewrite feedback that successfully re-runs Primary.
      2. Thumbs-up / thumbs-down ratings on Primary assistant results.

    Workbook write failures must never fail the user's request or rating.
    """

    def __init__(
        self,
        *,
        workbook_path: Optional[Path] = None,
    ) -> None:
        project_root = Path(__file__).resolve().parents[1]
        self.workbook_path = Path(
            workbook_path
            or project_root / "data" / "New Flagged FAQ.xlsx"
        )
        self.sheet_name = "Flagged Feedback"
        self._lock = asyncio.Lock()

    async def record_feedback_candidate(
        self,
        *,
        report_name: str,
        original_user_question: str,
        previous_normalized_question: str,
        feedback_input: str,
        normalized_feedback: str,
        new_normalized_question: str,
        variables: list[dict[str, Any]],
        variable_names: list[str],
        variables_key: str,
        sql_template: Optional[str],
        executed_sql: Optional[str],
        result_count: int,
        chat_id: Optional[str],
        source_assistant_message_id: Optional[str],
        workflow_id: Optional[str],
    ) -> dict[str, Any]:
        candidate_key = self._candidate_key(
            report_name=report_name,
            normalized_question=new_normalized_question,
            variables_key=variables_key,
            sql_template=sql_template,
        )

        row_payload = {
            "Candidate Key": candidate_key,
            "Created At UTC": self._utc_iso(),
            "Feedback Type": "modify",
            "Thumb Rating": "",
            "Review Priority": "medium",
            "Report Name": report_name,
            "Original User Question": original_user_question,
            "Previous Normalized Question": previous_normalized_question,
            "Feedback Input": feedback_input,
            "Normalized Feedback": normalized_feedback,
            "New Normalized Question": new_normalized_question,
            "Normalized Question": new_normalized_question,
            "Variable Names": ",".join(variable_names or []),
            "Variables Key": variables_key,
            "Variables JSON": self._json(variables or []),
            "Action": "",
            "Selected Example": "",
            "FAISS Examples Used": "",
            "SQL Template": sql_template,
            "Executed SQL": executed_sql,
            "Query Executed": bool(executed_sql),
            "Result Count": int(result_count or 0),
            "Results JSON": "",
            "Assistant Answer": "",
            "Request Mode": "feedback",
            "Chat ID": chat_id,
            "Source User Message ID": "",
            "Source Assistant Message ID": source_assistant_message_id,
            "Workflow ID": workflow_id,
            "Diagnostics JSON": "",
            "Review Status": "pending",
            "Display Question": "",
            "Review Notes": "",
        }

        async with self._lock:
            return await asyncio.to_thread(
                self._write_named_row,
                row_payload,
                False,
            )

    async def record_rating_candidate(
        self,
        *,
        rating: str,
        assistant_message: dict[str, Any],
        source_user_message: Optional[dict[str, Any]] = None,
    ) -> dict[str, Any]:
        """
        Record the exact Primary result that received a thumbs rating.

        One row is kept per assistant message. If the user changes a rating
        from up to down (or vice versa), the existing row is updated instead
        of creating a duplicate.
        """

        rating = str(rating or "").strip().lower()
        if rating not in {"up", "down"}:
            raise ValueError("Rating must be up or down.")

        workflow_used = str(
            assistant_message.get("workflow_used") or ""
        ).strip().lower()

        # The flagged FAQ workbook is intended to improve data-query SQL.
        # Do not fill it with greeting/help/metadata ratings.
        if workflow_used != "primary":
            return {
                "stored": False,
                "duplicate": False,
                "skipped": True,
                "reason": "not_primary_data_result",
                "path": str(self.workbook_path),
            }

        assistant_message_id = str(
            assistant_message.get("message_id") or ""
        ).strip()

        if not assistant_message_id:
            return {
                "stored": False,
                "duplicate": False,
                "skipped": True,
                "reason": "assistant_message_id_missing",
                "path": str(self.workbook_path),
            }

        candidate_key = hashlib.sha256(
            f"rating|{assistant_message_id}".encode("utf-8")
        ).hexdigest()

        variables = list(assistant_message.get("variables") or [])
        variable_names = list(
            assistant_message.get("variable_names") or []
        )
        if not variable_names:
            variable_names = self._input_variable_names(variables)

        variables_key = ",".join(sorted(variable_names))
        feedback_reexecution = assistant_message.get("feedback_reexecution") or {}

        normalized_question = str(
            assistant_message.get("resolved_normalized_input")
            or assistant_message.get("normalized_input")
            or ""
        ).strip()

        source_user_content = str(
            (source_user_message or {}).get("content") or ""
        ).strip()

        request_mode = str(
            assistant_message.get("request_mode") or ""
        ).strip().lower()

        # For a result created by Modify/feedback, the source user message is
        # the feedback text itself.  The original question being modified is
        # preserved in feedback_reexecution, so keep both parts of the chain.
        if request_mode == "feedback":
            original_user_question = str(
                feedback_reexecution.get("original_user_question") or ""
            ).strip()
            feedback_input = str(
                assistant_message.get("original_feedback_input")
                or feedback_reexecution.get("feedback_input")
                or source_user_content
            ).strip()
            normalized_feedback = str(
                assistant_message.get("normalized_feedback_input")
                or feedback_reexecution.get("normalized_feedback_input")
                or ""
            ).strip()
        else:
            original_user_question = source_user_content
            feedback_input = ""
            normalized_feedback = ""

        feedback_type = (
            "thumbs_up" if rating == "up" else "thumbs_down"
        )

        row_payload = {
            "Candidate Key": candidate_key,
            "Created At UTC": self._utc_iso(),
            "Feedback Type": feedback_type,
            "Thumb Rating": rating,
            "Review Priority": "normal" if rating == "up" else "high",
            "Report Name": assistant_message.get("report_name"),
            "Original User Question": original_user_question,
            "Previous Normalized Question": feedback_reexecution.get(
                "previous_normalized_question"
            ),
            "Feedback Input": feedback_input,
            "Normalized Feedback": normalized_feedback,
            "New Normalized Question": normalized_question,
            "Normalized Question": normalized_question,
            "Variable Names": ",".join(variable_names),
            "Variables Key": variables_key,
            "Variables JSON": self._json(variables),
            "Action": assistant_message.get("action"),
            "Selected Example": assistant_message.get("selected_example"),
            "FAISS Examples Used": assistant_message.get("faiss_examples_used"),
            "SQL Template": assistant_message.get("sql_template"),
            "Executed SQL": assistant_message.get("sql"),
            "Query Executed": bool(assistant_message.get("query_executed")),
            "Result Count": int(assistant_message.get("result_count") or 0),
            "Results JSON": self._json(assistant_message.get("table_rows") or []),
            "Assistant Answer": assistant_message.get("content"),
            "Request Mode": assistant_message.get("request_mode"),
            "Chat ID": assistant_message.get("chat_id"),
            "Source User Message ID": assistant_message.get(
                "source_user_message_id"
            ),
            "Source Assistant Message ID": assistant_message_id,
            "Workflow ID": assistant_message.get("workflow_id"),
            "Diagnostics JSON": self._json(feedback_reexecution),
            "Review Status": "pending",
            "Display Question": "",
            "Review Notes": "",
        }

        async with self._lock:
            return await asyncio.to_thread(
                self._write_named_row,
                row_payload,
                True,
            )

    def _write_named_row(
        self,
        payload: dict[str, Any],
        update_existing: bool,
    ) -> dict[str, Any]:
        self.workbook_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        if self.workbook_path.exists():
            workbook = load_workbook(self.workbook_path)
        else:
            workbook = Workbook()
            workbook.active.title = self.sheet_name

        if self.sheet_name in workbook.sheetnames:
            worksheet = workbook[self.sheet_name]
        else:
            worksheet = workbook.create_sheet(self.sheet_name)

        header_map = self._ensure_columns(worksheet)
        candidate_key = str(payload.get("Candidate Key") or "")
        existing_row = self._find_candidate_row(
            worksheet,
            header_map,
            candidate_key,
        )

        if existing_row is not None and not update_existing:
            workbook.close()
            return {
                "stored": False,
                "duplicate": True,
                "candidate_key": candidate_key,
                "path": str(self.workbook_path),
            }

        row_number = existing_row or (worksheet.max_row + 1)

        for column_name, value in payload.items():
            column_index = header_map.get(column_name)
            if column_index is None:
                continue
            worksheet.cell(
                row=row_number,
                column=column_index,
                value=self._excel_safe(value),
            )

        workbook.save(self.workbook_path)
        workbook.close()

        return {
            "stored": True,
            "duplicate": existing_row is not None,
            "updated": existing_row is not None,
            "candidate_key": candidate_key,
            "path": str(self.workbook_path),
        }

    @staticmethod
    def _ensure_columns(worksheet) -> dict[str, int]:
        existing_headers: dict[str, int] = {}

        for column_index in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row=1, column=column_index).value
            if value is not None and str(value).strip():
                existing_headers[str(value).strip()] = column_index

        if not existing_headers:
            for column_index, column_name in enumerate(
                DEFAULT_COLUMNS,
                start=1,
            ):
                worksheet.cell(
                    row=1,
                    column=column_index,
                    value=column_name,
                )
                existing_headers[column_name] = column_index
            worksheet.freeze_panes = "A2"
            return existing_headers

        next_column = worksheet.max_column + 1
        for column_name in DEFAULT_COLUMNS:
            if column_name in existing_headers:
                continue
            worksheet.cell(
                row=1,
                column=next_column,
                value=column_name,
            )
            existing_headers[column_name] = next_column
            next_column += 1

        worksheet.freeze_panes = "A2"
        return existing_headers

    @staticmethod
    def _find_candidate_row(
        worksheet,
        header_map: dict[str, int],
        candidate_key: str,
    ) -> Optional[int]:
        if not candidate_key:
            return None

        key_column = header_map.get("Candidate Key")
        if key_column is None:
            return None

        for row_number in range(2, worksheet.max_row + 1):
            value = worksheet.cell(
                row=row_number,
                column=key_column,
            ).value
            if str(value or "") == candidate_key:
                return row_number

        return None

    @staticmethod
    def _input_variable_names(
        variables: list[dict[str, Any]],
    ) -> list[str]:
        names: list[str] = []
        seen: set[str] = set()

        for item in variables:
            if not isinstance(item, dict):
                continue
            if str(item.get("source") or "input").lower() != "input":
                continue
            name = str(item.get("name") or "").strip().lower()
            if not name or name in seen:
                continue
            seen.add(name)
            names.append(name)

        return names

    @staticmethod
    def _json(value: Any) -> str:
        text = json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )
        # Excel cells have a 32,767-character limit. Keep a useful snapshot
        # rather than allowing workbook save to fail on very large results.
        if len(text) > 32000:
            return text[:31950] + "...<truncated>"
        return text

    @staticmethod
    def _excel_safe(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (dict, list, tuple, set)):
            return FlaggedFaqService._json(value)
        if isinstance(value, str) and len(value) > 32000:
            return value[:31950] + "...<truncated>"
        return value

    @staticmethod
    def _utc_iso() -> str:
        return datetime.now(timezone.utc).isoformat()

    @staticmethod
    def _candidate_key(
        *,
        report_name: str,
        normalized_question: str,
        variables_key: str,
        sql_template: Optional[str],
    ) -> str:
        raw = "|".join(
            [
                str(report_name or "").strip().lower(),
                str(normalized_question or "").strip().lower(),
                str(variables_key or "").strip().lower(),
                str(sql_template or "").strip(),
            ]
        )
        return hashlib.sha256(
            raw.encode("utf-8")
        ).hexdigest()

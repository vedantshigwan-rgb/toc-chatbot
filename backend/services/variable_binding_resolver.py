from __future__ import annotations

import os
import re
from collections import Counter
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook

from sqlglot import exp, parse_one
from sqlglot.errors import ParseError

from core.models import (
    SQLVariableBinding,
    VariableBindingIssue,
    VariableBindingResolutionResult,
)


PLACEHOLDER_PATTERN = re.compile(
    r"\{([A-Za-z_][A-Za-z0-9_]*)\}"
)

MARKER_PREFIX = "__TOC_VAR_"
MARKER_SUFFIX = "__"

MARKER_PATTERN = re.compile(
    r"__TOC_VAR_([A-Za-z_][A-Za-z0-9_]*)__"
)


PREDICATE_NODE_NAMES = {
    "Like",
    "ILike",
    "EQ",
    "NEQ",
    "GT",
    "GTE",
    "LT",
    "LTE",
    "In",
    "Between",
}


OPERATOR_NAMES = {
    "Like": "LIKE",
    "ILike": "ILIKE",
    "EQ": "=",
    "NEQ": "!=",
    "GT": ">",
    "GTE": ">=",
    "LT": "<",
    "LTE": "<=",
    "In": "IN",
    "Between": "BETWEEN",
}




DEFAULT_VARIABLE_MAPPING_WORKBOOK = (
    Path(__file__).resolve().parents[1]
    / "redis mapping.xlsx"
)

REQUIRED_MAPPING_HEADERS = {
    "report name",
    "variables",
    "column names",
    "table name",
    "storage_type",
}


def _normalize_mapping_text(
    value: Any,
) -> str:
    return re.sub(
        r"\s+",
        " ",
        str(value or "").strip(),
    )


def _normalize_mapping_key(
    value: Any,
) -> str:
    return (
        _normalize_mapping_text(value)
        .casefold()
    )


def _resolve_mapping_workbook_path(
    mapping_workbook_path: Optional[str],
) -> Path:
    configured_path = (
        mapping_workbook_path
        or os.getenv(
            "VARIABLE_MAPPING_WORKBOOK"
        )
    )

    if configured_path:
        return Path(
            configured_path
        ).expanduser().resolve()

    return (
        DEFAULT_VARIABLE_MAPPING_WORKBOOK
        .resolve()
    )


@lru_cache(maxsize=8)
def _load_variable_mapping_index(
    workbook_path: str,
    workbook_mtime_ns: int,
) -> Dict[
    Tuple[str, str],
    Tuple[Dict[str, str], ...],
]:
    """
    Load the mapping workbook and create:

        (normalized report, normalized variable)
            -> allowed mapping rows

    workbook_mtime_ns is part of the cache key. When the
    workbook changes, it is automatically reloaded.
    """

    workbook = load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
    )

    try:
        if not workbook.sheetnames:
            raise ValueError(
                "The variable mapping workbook "
                "has no worksheets."
            )

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        headers = [
            _normalize_mapping_key(
                cell.value
            )
            for cell in worksheet[1]
        ]

        header_indexes = {
            header: index
            for index, header
            in enumerate(headers)
            if header
        }

        missing_headers = sorted(
            REQUIRED_MAPPING_HEADERS
            - set(header_indexes)
        )

        if missing_headers:
            raise ValueError(
                "The variable mapping workbook "
                "is missing required headers: "
                + ", ".join(
                    missing_headers
                )
            )

        mapping_index: Dict[
            Tuple[str, str],
            List[Dict[str, str]],
        ] = {}

        seen_rows = set()

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            report_name = (
                _normalize_mapping_text(
                    row[
                        header_indexes[
                            "report name"
                        ]
                    ]
                )
            )

            variable_name = (
                _normalize_mapping_text(
                    row[
                        header_indexes[
                            "variables"
                        ]
                    ]
                )
            )

            column_name = (
                _normalize_mapping_text(
                    row[
                        header_indexes[
                            "column names"
                        ]
                    ]
                )
            )

            table_name = (
                _normalize_mapping_text(
                    row[
                        header_indexes[
                            "table name"
                        ]
                    ]
                )
            )

            storage_type = (
                _normalize_mapping_text(
                    row[
                        header_indexes[
                            "storage_type"
                        ]
                    ]
                )
            )

            if not (
                report_name
                and variable_name
                and column_name
                and table_name
            ):
                continue

            mapping_key = (
                _normalize_mapping_key(
                    report_name
                ),
                _normalize_mapping_key(
                    variable_name
                ),
            )

            row_key = (
                mapping_key,
                _normalize_mapping_key(
                    column_name
                ),
                _normalize_mapping_key(
                    table_name
                ),
                _normalize_mapping_key(
                    storage_type
                ),
            )

            if row_key in seen_rows:
                continue

            seen_rows.add(row_key)

            mapping_index.setdefault(
                mapping_key,
                [],
            ).append(
                {
                    "report_name": (
                        report_name
                    ),
                    "variable_name": (
                        variable_name
                    ),
                    "column_name": (
                        column_name
                    ),
                    "table_name": (
                        table_name
                    ),
                    "storage_type": (
                        storage_type
                    ),
                }
            )

        return {
            key: tuple(value)
            for key, value
            in mapping_index.items()
        }

    finally:
        workbook.close()


def _get_variable_mapping_index(
    mapping_workbook_path: Optional[str],
) -> Tuple[
    Path,
    Dict[
        Tuple[str, str],
        Tuple[Dict[str, str], ...],
    ],
]:
    workbook_path = (
        _resolve_mapping_workbook_path(
            mapping_workbook_path
        )
    )

    if not workbook_path.is_file():
        raise FileNotFoundError(
            "Variable mapping workbook "
            "not found: "
            f"{workbook_path}"
        )

    mapping_index = (
        _load_variable_mapping_index(
            str(workbook_path),
            workbook_path.stat().st_mtime_ns,
        )
    )

    return (
        workbook_path,
        mapping_index,
    )


def load_variable_mapping_rows(
    mapping_workbook_path: Optional[str] = None,
) -> Tuple[Path, List[Dict[str, str]]]:
    """
    Load every row of the variable mapping workbook as a flat list.

    Reuses the exact same cached loader as SQL-time fallback
    validation (`_get_variable_mapping_index`) - never a second,
    independent parser - so deterministic governance auditing
    (Step 38) always sees the identical mapping data the runtime
    binder itself would see, including automatic reload when the
    workbook file changes.

    Raises FileNotFoundError when the workbook does not exist, and
    ValueError when its structure is invalid (missing worksheet /
    required headers) - the same exceptions the runtime path raises.
    """

    _, mapping_index = _get_variable_mapping_index(
        mapping_workbook_path
    )

    rows: List[Dict[str, str]] = [
        dict(row)
        for candidates in mapping_index.values()
        for row in candidates
    ]

    workbook_path = _resolve_mapping_workbook_path(
        mapping_workbook_path
    )

    return workbook_path, rows


# def _qualified_mapping_table_name(
#     *,
#     schema_name: Optional[str],
#     table_name: str,
# ) -> str:
#     if "." in table_name:
#         return table_name

#     if schema_name:
#         return (
#             f"{schema_name}."
#             f"{table_name}"
#         )

#     return table_name


# def _apply_authoritative_mappings(
#     *,
#     bindings: List[
#         SQLVariableBinding
#     ],
#     placeholder_names: List[str],
#     report_name: str,
#     mapping_workbook_path: Optional[str],
#     issues: List[
#         VariableBindingIssue
#     ],
# ) -> Tuple[
#     List[SQLVariableBinding],
#     Set[str],
#     Set[str],
#     Dict[str, Any],
# ]:
#     """
#     Validate parsed SQL columns against redis mapping.xlsx.

#     The SQL AST determines which column the SQL currently uses.
#     The workbook determines whether that variable-to-column
#     relationship is allowed.

#     The workbook table becomes the lookup table used by the
#     no-data value validator.
#     """

#     diagnostics: Dict[str, Any] = {
#         "enabled": True,
#         "report_name": report_name,
#         "workbook_path": None,
#         "status": "not_loaded",
#     }

#     try:
#         (
#             workbook_path,
#             mapping_index,
#         ) = _get_variable_mapping_index(
#             mapping_workbook_path
#         )

#     except Exception as error:
#         diagnostics.update(
#             {
#                 "status": "error",
#                 "message": str(error),
#                 "exception_type": (
#                     type(error).__name__
#                 ),
#             }
#         )

#         issues.append(
#             VariableBindingIssue(
#                 code=(
#                     "variable_mapping_"
#                     "workbook_error"
#                 ),
#                 message=(
#                     "The authoritative variable "
#                     "mapping workbook could not "
#                     "be loaded."
#                 ),
#                 details=diagnostics,
#             )
#         )

#         return (
#             bindings,
#             set(placeholder_names),
#             set(),
#             diagnostics,
#         )

#     diagnostics.update(
#         {
#             "status": "loaded",
#             "workbook_path": str(
#                 workbook_path
#             ),
#             "mapping_key_count": (
#                 len(mapping_index)
#             ),
#         }
#     )

#     mapped_bindings: List[
#         SQLVariableBinding
#     ] = []

#     unresolved_variables: Set[str] = (
#         set()
#     )

#     ambiguous_variables: Set[str] = (
#         set()
#     )

#     emitted_issues = set()

#     for binding in bindings:
#         variable_name = (
#             binding.variable_name
#         )

#         mapping_key = (
#             _normalize_mapping_key(
#                 report_name
#             ),
#             _normalize_mapping_key(
#                 variable_name
#             ),
#         )

#         candidates = list(
#             mapping_index.get(
#                 mapping_key,
#                 (),
#             )
#         )

#         allowed_mappings = [
#             {
#                 "column_name": (
#                     candidate[
#                         "column_name"
#                     ]
#                 ),
#                 "table_name": (
#                     candidate[
#                         "table_name"
#                     ]
#                 ),
#                 "storage_type": (
#                     candidate[
#                         "storage_type"
#                     ]
#                 ),
#             }
#             for candidate in candidates
#         ]

#         if not candidates:
#             unresolved_variables.add(
#                 variable_name
#             )

#             issue_key = (
#                 "mapping_not_found",
#                 _normalize_mapping_key(
#                     variable_name
#                 ),
#             )

#             if issue_key not in emitted_issues:
#                 emitted_issues.add(
#                     issue_key
#                 )

#                 issues.append(
#                     VariableBindingIssue(
#                         code=(
#                             "variable_mapping_"
#                             "not_found"
#                         ),
#                         message=(
#                             "No authoritative mapping "
#                             "was found for this report "
#                             "variable."
#                         ),
#                         variable_name=(
#                             variable_name
#                         ),
#                         details={
#                             "report_name": (
#                                 report_name
#                             ),
#                         },
#                     )
#                 )

#             raw_binding = (
#                 binding.model_dump()
#             )

#             raw_binding.update(
#                 {
#                     "mapping_validated": (
#                         False
#                     ),
#                     "mapping_source": (
#                         "redis mapping.xlsx"
#                     ),
#                     "mapping_candidates": [],
#                     "sql_table_name": (
#                         binding.table_name
#                     ),
#                     "sql_column_name": (
#                         binding.column_name
#                     ),
#                 }
#             )

#             mapped_bindings.append(
#                 SQLVariableBinding(
#                     **raw_binding
#                 )
#             )

#             continue

#         column_matches = [
#             candidate
#             for candidate in candidates
#             if (
#                 _normalize_mapping_key(
#                     candidate[
#                         "column_name"
#                     ]
#                 )
#                 == _normalize_mapping_key(
#                     binding.column_name
#                 )
#             )
#         ]

#         if not column_matches:
#             unresolved_variables.add(
#                 variable_name
#             )

#             issue_key = (
#                 "column_mismatch",
#                 _normalize_mapping_key(
#                     variable_name
#                 ),
#                 _normalize_mapping_key(
#                     binding.column_name
#                 ),
#             )

#             if issue_key not in emitted_issues:
#                 emitted_issues.add(
#                     issue_key
#                 )

#                 issues.append(
#                     VariableBindingIssue(
#                         code=(
#                             "variable_column_"
#                             "mapping_mismatch"
#                         ),
#                         message=(
#                             "The SQL predicate column "
#                             "does not match an "
#                             "authoritative mapping for "
#                             "the report variable."
#                         ),
#                         variable_name=(
#                             variable_name
#                         ),
#                         details={
#                             "report_name": (
#                                 report_name
#                             ),
#                             "sql_table_name": (
#                                 binding.table_name
#                             ),
#                             "sql_column_name": (
#                                 binding.column_name
#                             ),
#                             "allowed_mappings": (
#                                 allowed_mappings
#                             ),
#                         },
#                     )
#                 )

#             raw_binding = (
#                 binding.model_dump()
#             )

#             raw_binding.update(
#                 {
#                     "mapping_validated": (
#                         False
#                     ),
#                     "mapping_source": (
#                         "redis mapping.xlsx"
#                     ),
#                     "mapping_candidates": (
#                         allowed_mappings
#                     ),
#                     "sql_table_name": (
#                         binding.table_name
#                     ),
#                     "sql_column_name": (
#                         binding.column_name
#                     ),
#                 }
#             )

#             mapped_bindings.append(
#                 SQLVariableBinding(
#                     **raw_binding
#                 )
#             )

#             continue

#         unique_matches = []

#         seen_targets = set()

#         for candidate in column_matches:
#             target_key = (
#                 _normalize_mapping_key(
#                     candidate[
#                         "table_name"
#                     ]
#                 ),
#                 _normalize_mapping_key(
#                     candidate[
#                         "column_name"
#                     ]
#                 ),
#             )

#             if target_key in seen_targets:
#                 continue

#             seen_targets.add(
#                 target_key
#             )

#             unique_matches.append(
#                 candidate
#             )

#         selected_mapping = None

#         if len(unique_matches) == 1:
#             selected_mapping = (
#                 unique_matches[0]
#             )

#         else:
#             sql_table_matches = [
#                 candidate
#                 for candidate
#                 in unique_matches
#                 if (
#                     _normalize_mapping_key(
#                         candidate[
#                             "table_name"
#                         ]
#                     )
#                     == _normalize_mapping_key(
#                         binding.table_name
#                     )
#                 )
#             ]

#             if len(sql_table_matches) == 1:
#                 selected_mapping = (
#                     sql_table_matches[0]
#                 )

#         if selected_mapping is None:
#             ambiguous_variables.add(
#                 variable_name
#             )

#             issue_key = (
#                 "mapping_ambiguous",
#                 _normalize_mapping_key(
#                     variable_name
#                 ),
#                 _normalize_mapping_key(
#                     binding.column_name
#                 ),
#             )

#             if issue_key not in emitted_issues:
#                 emitted_issues.add(
#                     issue_key
#                 )

#                 issues.append(
#                     VariableBindingIssue(
#                         code=(
#                             "variable_mapping_"
#                             "ambiguous"
#                         ),
#                         message=(
#                             "More than one "
#                             "authoritative mapping "
#                             "matches the SQL column."
#                         ),
#                         variable_name=(
#                             variable_name
#                         ),
#                         details={
#                             "report_name": (
#                                 report_name
#                             ),
#                             "sql_table_name": (
#                                 binding.table_name
#                             ),
#                             "sql_column_name": (
#                                 binding.column_name
#                             ),
#                             "matching_mappings": (
#                                 unique_matches
#                             ),
#                         },
#                     )
#                 )

#             mapped_bindings.append(
#                 binding
#             )

#             continue

#         mapped_table = (
#             selected_mapping[
#                 "table_name"
#             ]
#         )

#         mapped_column = (
#             selected_mapping[
#                 "column_name"
#             ]
#         )

#         mapped_binding = (
#             binding.model_dump()
#         )

#         mapped_binding.update(
#             {
#                 # Authoritative value-validation
#                 # target from the workbook.
#                 "table_name": mapped_table,
#                 "qualified_table_name": (
#                     _qualified_mapping_table_name(
#                         schema_name=(
#                             binding.schema_name
#                         ),
#                         table_name=(
#                             mapped_table
#                         ),
#                     )
#                 ),
#                 "column_name": mapped_column,

#                 # Diagnostics showing what was
#                 # parsed from the generated SQL.
#                 "sql_table_name": (
#                     binding.table_name
#                 ),
#                 "sql_column_name": (
#                     binding.column_name
#                 ),

#                 "mapping_validated": True,
#                 "mapping_source": (
#                     "redis mapping.xlsx"
#                 ),
#                 "mapping_report_name": (
#                     selected_mapping[
#                         "report_name"
#                     ]
#                 ),
#                 "storage_type": (
#                     selected_mapping[
#                         "storage_type"
#                     ]
#                 ),
#                 "mapping_candidates": (
#                     allowed_mappings
#                 ),
#             }
#         )

#         mapped_bindings.append(
#             SQLVariableBinding(
#                 **mapped_binding
#             )
#         )

#     return (
#         mapped_bindings,
#         unresolved_variables,
#         ambiguous_variables,
#         diagnostics,
#     )




def _normalize_identifier(
    value: Any,
) -> str:
    return str(
        value or ""
    ).strip().casefold()


def _mask_placeholders(
    sql_template: str,
) -> str:
    """
    Replace braces with SQL-parseable marker names.

    Examples:

        '%{buyer_name}%'
            -> '%__TOC_VAR_buyer_name__%'

        QUANTITY > {minimum_qty}
            -> QUANTITY > __TOC_VAR_minimum_qty__
    """

    return PLACEHOLDER_PATTERN.sub(
        lambda match: (
            f"{MARKER_PREFIX}"
            f"{match.group(1)}"
            f"{MARKER_SUFFIX}"
        ),
        sql_template,
    )


def _variables_in_expression(
    expression: Optional[exp.Expression],
    *,
    dialect: str,
) -> Set[str]:
    if expression is None:
        return set()

    try:
        expression_sql = expression.sql(
            dialect=dialect
        )
    except Exception:
        expression_sql = str(expression)

    return set(
        MARKER_PATTERN.findall(
            expression_sql
        )
    )


def _expression_sql(
    expression: Optional[exp.Expression],
    *,
    dialect: str,
) -> str:
    if expression is None:
        return ""

    try:
        return expression.sql(
            dialect=dialect
        )
    except Exception:
        return str(expression)


def _is_variable_marker(
    value: str,
) -> bool:
    return bool(
        MARKER_PATTERN.fullmatch(
            str(value or "").strip()
        )
    )


def _collect_cte_names(
    statement: exp.Expression,
) -> Set[str]:
    cte_names: Set[str] = set()

    for cte in statement.find_all(
        exp.CTE
    ):
        cte_name = str(
            cte.alias_or_name or ""
        ).strip()

        if cte_name:
            cte_names.add(
                _normalize_identifier(
                    cte_name
                )
            )

    return cte_names


def _collect_physical_tables(
    statement: exp.Expression,
) -> List[Dict[str, Optional[str]]]:
    """
    Extract physical tables while excluding CTE aliases.
    """

    cte_names = _collect_cte_names(
        statement
    )

    tables: List[
        Dict[str, Optional[str]]
    ] = []

    seen = set()

    for table in statement.find_all(
        exp.Table
    ):
        table_name = str(
            table.name or ""
        ).strip()

        schema_name = str(
            table.db or ""
        ).strip() or None

        catalog_name = str(
            table.catalog or ""
        ).strip() or None

        alias = str(
            table.alias or ""
        ).strip() or None

        if not table_name:
            continue

        if (
            schema_name is None
            and catalog_name is None
            and _normalize_identifier(
                table_name
            )
            in cte_names
        ):
            continue

        qualified_parts = [
            catalog_name,
            schema_name,
            table_name,
        ]

        qualified_name = ".".join(
            part
            for part in qualified_parts
            if part
        )

        deduplication_key = (
            _normalize_identifier(
                qualified_name
            ),
            _normalize_identifier(alias),
        )

        if deduplication_key in seen:
            continue

        seen.add(
            deduplication_key
        )

        tables.append(
            {
                "catalog_name": catalog_name,
                "schema_name": schema_name,
                "table_name": table_name,
                "qualified_name": (
                    qualified_name
                ),
                "alias": alias,
            }
        )

    return tables


def _build_table_lookup(
    tables: List[
        Dict[str, Optional[str]]
    ],
) -> Dict[
    str,
    List[Dict[str, Optional[str]]],
]:
    lookup: Dict[
        str,
        List[Dict[str, Optional[str]]],
    ] = {}

    for table in tables:
        lookup_keys = {
            _normalize_identifier(
                table.get("table_name")
            ),
            _normalize_identifier(
                table.get("alias")
            ),
        }

        for lookup_key in lookup_keys:
            if not lookup_key:
                continue

            existing_tables = lookup.setdefault(
                lookup_key,
                [],
            )

            if table not in existing_tables:
                existing_tables.append(
                    table
                )

    return lookup


def _collect_columns(
    expression: Optional[exp.Expression],
) -> List[exp.Column]:
    if expression is None:
        return []

    columns: List[exp.Column] = []
    seen = set()

    for column in expression.find_all(
        exp.Column
    ):
        column_name = str(
            column.name or ""
        ).strip()

        table_alias = str(
            column.table or ""
        ).strip()

        if not column_name:
            continue

        if _is_variable_marker(
            column_name
        ):
            continue

        key = (
            _normalize_identifier(
                table_alias
            ),
            _normalize_identifier(
                column_name
            ),
        )

        if key in seen:
            continue

        seen.add(key)
        columns.append(column)

    return columns


def _predicate_sides(
    predicate: exp.Expression,
    *,
    dialect: str,
) -> Tuple[
    Optional[exp.Expression],
    str,
    Set[str],
]:
    """
    Return:

        candidate column expression,
        value-side SQL,
        variables found in the value side.
    """

    node_name = type(
        predicate
    ).__name__

    if node_name == "In":
        candidate_expression = (
            predicate.args.get("this")
        )

        value_expressions = list(
            predicate.args.get(
                "expressions"
            )
            or []
        )

        query_expression = (
            predicate.args.get("query")
        )

        if query_expression is not None:
            value_expressions.append(
                query_expression
            )

        value_sql = " ".join(
            _expression_sql(
                expression,
                dialect=dialect,
            )
            for expression in value_expressions
        )

        variables = set(
            MARKER_PATTERN.findall(
                value_sql
            )
        )

        return (
            candidate_expression,
            value_sql,
            variables,
        )

    if node_name == "Between":
        candidate_expression = (
            predicate.args.get("this")
        )

        low_expression = (
            predicate.args.get("low")
        )

        high_expression = (
            predicate.args.get("high")
        )

        value_sql = " ".join(
            [
                _expression_sql(
                    low_expression,
                    dialect=dialect,
                ),
                _expression_sql(
                    high_expression,
                    dialect=dialect,
                ),
            ]
        )

        variables = set(
            MARKER_PATTERN.findall(
                value_sql
            )
        )

        return (
            candidate_expression,
            value_sql,
            variables,
        )

    left_expression = (
        predicate.args.get("this")
    )

    right_expression = (
        predicate.args.get("expression")
    )

    left_variables = (
        _variables_in_expression(
            left_expression,
            dialect=dialect,
        )
    )

    right_variables = (
        _variables_in_expression(
            right_expression,
            dialect=dialect,
        )
    )

    if right_variables and not left_variables:
        return (
            left_expression,
            _expression_sql(
                right_expression,
                dialect=dialect,
            ),
            right_variables,
        )

    if left_variables and not right_variables:
        return (
            right_expression,
            _expression_sql(
                left_expression,
                dialect=dialect,
            ),
            left_variables,
        )

    all_variables = (
        left_variables
        | right_variables
    )

    return (
        predicate,
        _expression_sql(
            predicate,
            dialect=dialect,
        ),
        all_variables,
    )


def _determine_match_type(
    *,
    operator: str,
    variable_name: str,
    value_sql: str,
) -> str:
    if operator == "BETWEEN":
        return "range"

    if operator == "IN":
        return "in"

    if operator not in {
        "LIKE",
        "ILIKE",
    }:
        return "exact"

    marker = (
        f"{MARKER_PREFIX}"
        f"{variable_name}"
        f"{MARKER_SUFFIX}"
    )

    marker_position = value_sql.find(
        marker
    )

    if marker_position < 0:
        return "contains"

    before_marker = value_sql[
        :marker_position
    ]

    after_marker = value_sql[
        marker_position
        + len(marker):
    ]

    has_left_percent = (
        before_marker.rstrip().endswith(
            "%"
        )
    )

    has_right_percent = (
        after_marker.lstrip().startswith(
            "%"
        )
    )

    if (
        has_left_percent
        and has_right_percent
    ):
        return "contains"

    if has_right_percent:
        return "starts_with"

    if has_left_percent:
        return "ends_with"

    return "exact"


def _resolve_table_for_column(
    *,
    column: exp.Column,
    tables: List[
        Dict[str, Optional[str]]
    ],
    table_lookup: Dict[
        str,
        List[Dict[str, Optional[str]]],
    ],
) -> Tuple[
    Optional[Dict[str, Optional[str]]],
    Optional[str],
]:
    table_alias = str(
        column.table or ""
    ).strip()

    if table_alias:
        candidates = table_lookup.get(
            _normalize_identifier(
                table_alias
            ),
            [],
        )

        if len(candidates) == 1:
            return candidates[0], None

        if not candidates:
            return (
                None,
                "column_table_alias_not_found",
            )

        return (
            None,
            "column_table_alias_ambiguous",
        )

    if len(tables) == 1:
        return tables[0], None

    if not tables:
        return (
            None,
            "physical_table_not_found",
        )

    return (
        None,
        "unqualified_column_with_multiple_tables",
    )


def resolve_variable_bindings(
    *,
    sql_template: str,
    dialect: str = "oracle",
) -> VariableBindingResolutionResult:
    """
    Resolve SQL placeholders to their physical table columns.

    The function does not query Oracle or Weaviate. It only
    analyzes the SQL template.
    """

    normalized_template = str(
        sql_template or ""
    ).strip()

    placeholder_names = sorted(
        set(
            PLACEHOLDER_PATTERN.findall(
                normalized_template
            )
        )
    )

    placeholder_counts = Counter(
        PLACEHOLDER_PATTERN.findall(
            normalized_template
        )
    )

    if not normalized_template:
        return VariableBindingResolutionResult(
            success=False,
            complete=False,
            dialect=dialect,
            placeholder_names=[],
            unresolved_variables=[],
            ambiguous_variables=[],
            issues=[
                VariableBindingIssue(
                    code="empty_sql_template",
                    message=(
                        "The SQL template is empty."
                    ),
                )
            ],
        )

    if not placeholder_names:
        return VariableBindingResolutionResult(
            success=True,
            complete=True,
            dialect=dialect,
            placeholder_names=[],
            bindings=[],
            unresolved_variables=[],
            ambiguous_variables=[],
            issues=[],
            masked_sql=normalized_template,
        )

    masked_sql = _mask_placeholders(
        normalized_template
    )

    try:
        statement = parse_one(
            masked_sql,
            read=dialect,
        )

    except ParseError as error:
        return VariableBindingResolutionResult(
            success=False,
            complete=False,
            dialect=dialect,
            placeholder_names=(
                placeholder_names
            ),
            bindings=[],
            unresolved_variables=(
                placeholder_names
            ),
            ambiguous_variables=[],
            issues=[
                VariableBindingIssue(
                    code="binding_sql_parse_error",
                    message=str(error),
                    details={
                        "exception_type": (
                            type(error).__name__
                        ),
                    },
                )
            ],
            masked_sql=masked_sql,
        )

    except Exception as error:
        return VariableBindingResolutionResult(
            success=False,
            complete=False,
            dialect=dialect,
            placeholder_names=(
                placeholder_names
            ),
            bindings=[],
            unresolved_variables=(
                placeholder_names
            ),
            ambiguous_variables=[],
            issues=[
                VariableBindingIssue(
                    code=(
                        "binding_resolver_failure"
                    ),
                    message=str(error),
                    details={
                        "exception_type": (
                            type(error).__name__
                        ),
                    },
                )
            ],
            masked_sql=masked_sql,
        )

    tables = _collect_physical_tables(
        statement
    )

    table_lookup = _build_table_lookup(
        tables
    )

    issues: List[
        VariableBindingIssue
    ] = []

    binding_data: Dict[
        Tuple[
            str,
            Optional[str],
            Optional[str],
            Optional[str],
            str,
            str,
            str,
        ],
        Dict[str, Any],
    ] = {}

    resolved_variables: Set[str] = set()

    for node in statement.walk():
        node_name = type(
            node
        ).__name__

        if (
            node_name
            not in PREDICATE_NODE_NAMES
        ):
            continue

        (
            candidate_expression,
            value_sql,
            predicate_variables,
        ) = _predicate_sides(
            node,
            dialect=dialect,
        )

        if not predicate_variables:
            continue

        columns = _collect_columns(
            candidate_expression
        )

        predicate_sql = _expression_sql(
            node,
            dialect=dialect,
        )

        if len(columns) != 1:
            issue_code = (
                "binding_column_not_found"
                if not columns
                else (
                    "multiple_binding_columns"
                )
            )

            for variable_name in sorted(
                predicate_variables
            ):
                issues.append(
                    VariableBindingIssue(
                        code=issue_code,
                        message=(
                            "The resolver could not "
                            "determine one unique "
                            "column for the variable."
                        ),
                        variable_name=(
                            variable_name
                        ),
                        details={
                            "predicate_sql": (
                                predicate_sql
                            ),
                            "column_candidates": [
                                column.sql(
                                    dialect=dialect
                                )
                                for column in columns
                            ],
                        },
                    )
                )

            continue

        column = columns[0]

        table_info, table_issue = (
            _resolve_table_for_column(
                column=column,
                tables=tables,
                table_lookup=table_lookup,
            )
        )

        if table_issue is not None:
            for variable_name in sorted(
                predicate_variables
            ):
                issues.append(
                    VariableBindingIssue(
                        code=table_issue,
                        message=(
                            "The resolver could not "
                            "determine one physical "
                            "table for the column."
                        ),
                        variable_name=(
                            variable_name
                        ),
                        details={
                            "column": column.sql(
                                dialect=dialect
                            ),
                            "predicate_sql": (
                                predicate_sql
                            ),
                            "physical_tables": (
                                tables
                            ),
                        },
                    )
                )

            continue

        column_name = str(
            column.name or ""
        ).strip()

        table_alias = str(
            column.table or ""
        ).strip() or None

        operator = OPERATOR_NAMES.get(
            node_name,
            node_name.upper(),
        )

        for variable_name in sorted(
            predicate_variables
        ):
            match_type = (
                _determine_match_type(
                    operator=operator,
                    variable_name=(
                        variable_name
                    ),
                    value_sql=value_sql,
                )
            )

            binding_key = (
                variable_name,
                table_info.get(
                    "schema_name"
                ),
                table_info.get(
                    "table_name"
                ),
                table_alias,
                column_name,
                operator,
                match_type,
            )

            if binding_key not in binding_data:
                binding_data[
                    binding_key
                ] = {
                    "variable_name": (
                        variable_name
                    ),
                    "schema_name": (
                        table_info.get(
                            "schema_name"
                        )
                    ),
                    "table_name": (
                        table_info.get(
                            "table_name"
                        )
                    ),
                    "qualified_table_name": (
                        table_info.get(
                            "qualified_name"
                        )
                    ),
                    "table_alias": (
                        table_alias
                    ),
                    "column_name": (
                        column_name
                    ),
                    "operator": operator,
                    "match_type": (
                        match_type
                    ),
                    "placeholder_count": 0,
                    "predicate_sql": (
                        predicate_sql
                    ),
                }

            binding_data[
                binding_key
            ]["placeholder_count"] += 1

            resolved_variables.add(
                variable_name
            )

    bindings = [
        SQLVariableBinding(**data)
        for data in binding_data.values()
    ]

    bindings.sort(
        key=lambda binding: (
            binding.variable_name,
            binding.qualified_table_name
            or "",
            binding.column_name,
        )
    )



    targets_by_variable: Dict[
        str,
        Set[
            Tuple[
                Optional[str],
                Optional[str],
                str,
            ]
        ],
    ] = {}

    for binding in bindings:
        targets_by_variable.setdefault(
            binding.variable_name,
            set(),
        ).add(
            (
                binding.schema_name,
                binding.table_name,
                binding.column_name,
            )
        )

    ambiguous_variables = sorted(
        variable_name
        for (
            variable_name,
            targets,
        ) in targets_by_variable.items()
        if len(targets) > 1
    )

    for variable_name in (
        ambiguous_variables
    ):
        issues.append(
            VariableBindingIssue(
                code=(
                    "variable_bound_to_"
                    "multiple_columns"
                ),
                message=(
                    "The same variable is bound "
                    "to multiple physical SQL "
                    "columns."
                ),
                variable_name=(
                    variable_name
                ),
                details={
                    "targets": [
                        {
                            "schema_name": (
                                target[0]
                            ),
                            "table_name": (
                                target[1]
                            ),
                            "column_name": (
                                target[2]
                            ),
                        }
                        for target in sorted(
                            targets_by_variable[
                                variable_name
                            ],
                            key=lambda target: (
                                str(target[0]),
                                str(target[1]),
                                str(target[2]),
                            ),
                        )
                    ],
                },
            )
        )

    unresolved_variables = sorted(
        set(placeholder_names)
        - resolved_variables
    )

    for variable_name in (
        unresolved_variables
    ):
        issues.append(
            VariableBindingIssue(
                code=(
                    "unresolved_variable_binding"
                ),
                message=(
                    "No column predicate was "
                    "found for the SQL variable."
                ),
                variable_name=(
                    variable_name
                ),
                details={
                    "placeholder_count": (
                        placeholder_counts[
                            variable_name
                        ]
                    ),
                },
            )
        )

    complete = (
        not unresolved_variables
        and not ambiguous_variables
    )

    return VariableBindingResolutionResult(
        success=True,
        complete=complete,
        dialect=dialect,
        placeholder_names=(
            placeholder_names
        ),
        bindings=bindings,
        unresolved_variables=(
            unresolved_variables
        ),
        ambiguous_variables=(
            ambiguous_variables
        ),
        issues=issues,
        masked_sql=masked_sql,
    )


def validate_bindings_with_mapping_fallback(
    *,
    binding_result: VariableBindingResolutionResult,
    report_name: str,
    mapping_workbook_path: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Validate already-resolved SQL variable bindings against
    redis mapping.xlsx as a FALLBACK source only.

    Important rules:

    1. The SQL binding is never overwritten.
    2. Mapping is checked only against the SAME SQL table.
    3. Missing mapping for a table is UNVERIFIED, not invalid.
    4. A conflict exists only when the workbook explicitly
       maps the same report + variable + table to another column.

    Oracle remains the primary source of truth.
    """

    diagnostics: Dict[str, Any] = {
        "source": "redis mapping.xlsx",
        "authoritative": False,
        "fallback_used": True,
        "report_name": str(
            report_name or ""
        ).strip(),
        "success": False,
        "complete": False,
        "conflict": False,
        "status": "not_started",
        "bindings": [],
    }

    try:
        (
            workbook_path,
            mapping_index,
        ) = _get_variable_mapping_index(
            mapping_workbook_path
        )

    except Exception as error:
        diagnostics.update(
            {
                "status": "unavailable",
                "success": False,
                "message": str(error),
                "exception_type": (
                    type(error).__name__
                ),
            }
        )

        return diagnostics

    diagnostics["workbook_path"] = str(
        workbook_path
    )

    normalized_report = (
        _normalize_mapping_key(
            report_name
        )
    )

    binding_results = []

    explicit_conflicts = []
    matched_count = 0

    for binding in binding_result.bindings:
        variable_name = str(
            binding.variable_name or ""
        ).strip()

        sql_table_name = str(
            binding.table_name or ""
        ).strip()

        sql_column_name = str(
            binding.column_name or ""
        ).strip()

        mapping_key = (
            normalized_report,
            _normalize_mapping_key(
                variable_name
            ),
        )

        candidates = list(
            mapping_index.get(
                mapping_key,
                (),
            )
        )

        # ----------------------------------------------------
        # Compare BASE table names.
        #
        # KPCL_REPORTING.TABLE_A
        # and
        # TABLE_A
        #
        # should be considered the same table.
        # ----------------------------------------------------

        sql_base_table = (
            sql_table_name
            .split(".")[-1]
            .strip()
            .strip('"')
        )

        same_table_candidates = [
            candidate
            for candidate in candidates
            if (
                _normalize_mapping_key(
                    str(
                        candidate.get(
                            "table_name",
                            "",
                        )
                    )
                    .split(".")[-1]
                    .strip()
                    .strip('"')
                )
                == _normalize_mapping_key(
                    sql_base_table
                )
            )
        ]

        # ----------------------------------------------------
        # No mapping for the variable at all.
        #
        # This is NOT a failure because the workbook can be
        # incomplete/stale.
        # ----------------------------------------------------

        if not candidates:
            binding_results.append(
                {
                    "variable_name": (
                        variable_name
                    ),
                    "sql_table_name": (
                        sql_table_name
                    ),
                    "sql_column_name": (
                        sql_column_name
                    ),
                    "status": (
                        "variable_not_mapped"
                    ),
                    "validated": None,
                    "message": (
                        "No fallback mapping exists "
                        "for this variable."
                    ),
                    "candidates": [],
                }
            )

            continue

        # ----------------------------------------------------
        # Variable exists in mapping, but not for this SQL table.
        #
        # Example:
        #
        # workbook:
        # buyer_name -> FULLKIT_ITEMWISE_INV.BUYER_NAME
        #
        # generated SQL:
        # KPC_INVENTORY_BUYER_REP.BUYER
        #
        # Do NOT call this invalid.
        # ----------------------------------------------------

        if not same_table_candidates:
            binding_results.append(
                {
                    "variable_name": (
                        variable_name
                    ),
                    "sql_table_name": (
                        sql_table_name
                    ),
                    "sql_column_name": (
                        sql_column_name
                    ),
                    "status": (
                        "table_not_mapped"
                    ),
                    "validated": None,
                    "message": (
                        "Fallback mappings exist "
                        "for this variable, but none "
                        "exist for the SQL table."
                    ),
                    "candidates": [
                        {
                            "table_name": (
                                candidate.get(
                                    "table_name"
                                )
                            ),
                            "column_name": (
                                candidate.get(
                                    "column_name"
                                )
                            ),
                            "storage_type": (
                                candidate.get(
                                    "storage_type"
                                )
                            ),
                        }
                        for candidate
                        in candidates
                    ],
                }
            )

            continue

        # ----------------------------------------------------
        # Same table exists in mapping.
        # Now compare the column.
        # ----------------------------------------------------

        matching_column_candidates = [
            candidate
            for candidate
            in same_table_candidates
            if (
                _normalize_mapping_key(
                    candidate.get(
                        "column_name"
                    )
                )
                == _normalize_mapping_key(
                    sql_column_name
                )
            )
        ]

        if matching_column_candidates:
            matched_count += 1

            binding_results.append(
                {
                    "variable_name": (
                        variable_name
                    ),
                    "sql_table_name": (
                        sql_table_name
                    ),
                    "sql_column_name": (
                        sql_column_name
                    ),
                    "status": "matched",
                    "validated": True,
                    "message": (
                        "Fallback mapping supports "
                        "the SQL variable binding."
                    ),
                    "candidates": [
                        {
                            "table_name": (
                                candidate.get(
                                    "table_name"
                                )
                            ),
                            "column_name": (
                                candidate.get(
                                    "column_name"
                                )
                            ),
                            "storage_type": (
                                candidate.get(
                                    "storage_type"
                                )
                            ),
                        }
                        for candidate
                        in same_table_candidates
                    ],
                }
            )

            continue

        # ----------------------------------------------------
        # This is the ONLY explicit mapping conflict:
        #
        # same report
        # same variable
        # same table
        # different column
        # ----------------------------------------------------

        conflict = {
            "variable_name": variable_name,
            "sql_table_name": (
                sql_table_name
            ),
            "sql_column_name": (
                sql_column_name
            ),
            "status": (
                "column_conflict"
            ),
            "validated": False,
            "message": (
                "Fallback mapping contains the "
                "same table for this variable, "
                "but maps it to a different "
                "column."
            ),
            "candidates": [
                {
                    "table_name": (
                        candidate.get(
                            "table_name"
                        )
                    ),
                    "column_name": (
                        candidate.get(
                            "column_name"
                        )
                    ),
                    "storage_type": (
                        candidate.get(
                            "storage_type"
                        )
                    ),
                }
                for candidate
                in same_table_candidates
            ],
        }

        explicit_conflicts.append(
            conflict
        )

        binding_results.append(
            conflict
        )

    diagnostics.update(
        {
            "success": True,
            "status": (
                "conflict"
                if explicit_conflicts
                else (
                    "validated"
                    if (
                        binding_results
                        and matched_count
                        == len(
                            binding_results
                        )
                    )
                    else "partially_verified"
                )
            ),
            "complete": (
                bool(binding_results)
                and matched_count
                == len(binding_results)
            ),
            "conflict": bool(
                explicit_conflicts
            ),
            "matched_count": (
                matched_count
            ),
            "binding_count": len(
                binding_results
            ),
            "bindings": (
                binding_results
            ),
            "conflicts": (
                explicit_conflicts
            ),
        }
    )

    return diagnostics